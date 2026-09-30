from copy import deepcopy
import csv
import tkinter as tk
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from price_review import ReviewSession, PRICE, CODE
from price_review_ui import ReviewWindow
from candidate_review_ui import CandidateReviewWindow
import test_price_review


def candidate_fixture():
    fixture = test_price_review.SessionTests()
    fixture.setUp()
    fixture.rows[0].update({PRICE: 11, "Target Quantity kg": 2, "Average Mode": "store_balanced",
                            "Market Price Level": "average", "BAP_MARKET_PRICE_PERCENT": .2,
                            "STORE_WEIGHT_CAP": 1, "ECONOMY_PERCENTILE": .25,
                            "Selected Stores": "rey, ribasmith, super99"})
    fixture.write_results()
    wb = load_workbook(fixture.results)
    ws = wb.create_sheet("Candidates")
    ws.append([CODE, "Store", "Store Name", "Status", "SKU", "Candidate Product", "Chosen Price",
               "Normalized kg", "Price / kg", "Reason", "Product URL"])
    for index, (store, status, price) in enumerate((
            ("rey", "ACCEPTED", 10), ("rey", "ACCEPTED", 20),
            ("ribasmith", "ACCEPTED", 40), ("ribasmith", "REJECTED", 60),
            ("ribasmith", "REJECTED", None))):
        ws.append(["001", store, store, status, str(index), f"Rice {index}", price,
                   1 if price else None, price, "matching product" if status == "ACCEPTED" else "size mismatch",
                   "https://example.com/product"])
    wb.save(fixture.results)
    wb.close()
    return fixture


class CandidateReviewTests(unittest.TestCase):
    def setUp(self):
        self.fixture = candidate_fixture()
        self.addCleanup(self.fixture.tearDown)
        self.session = self.fixture.session()

    def test_preview_uses_original_formula_and_preserves_inputs(self):
        session = self.session
        before = deepcopy(session.items["001"])
        self.assertAlmostEqual(session.preview_candidates("001", {})["row"][PRICE], 11)
        self.assertAlmostEqual(session.preview_candidates("001", {"1": False})["row"][PRICE], 10)
        self.assertAlmostEqual(session.preview_candidates("001", {"3": True})["row"][PRICE], 13)
        self.assertEqual(session.items["001"], before)
        self.assertFalse(session.path.exists())

    def test_apply_reopen_approve_export_and_undo(self):
        session = self.session
        session.decide(["001"], "accepted")
        session.apply_candidates("001", {"3": True})
        self.assertEqual(session.items["001"]["decision"], "pending")
        self.assertIsNone(session.items["001"]["price"])
        reopened = self.fixture.session()
        self.assertEqual(reopened.items["001"]["row"][PRICE], 13)
        self.assertEqual(reopened.items["001"]["counts"], {"rey": 2, "ribasmith": 2})
        self.assertTrue(reopened.candidate_accepted("001", 3))
        self.assertEqual(reopened.candidates["001"][3]["Status"], "REJECTED")
        reopened.decide(["001"], "accepted")
        destination = self.fixture.root / "export.csv"
        reopened.export(destination)
        with destination.open(encoding="utf-8") as stream:
            self.assertEqual(list(csv.DictReader(stream))[0]["UnitPrice"], "13.00")
        session.undo()
        self.assertEqual(session.items["001"]["decision"], "accepted")
        self.assertEqual(session.items["001"]["price"], 11)
        self.assertFalse(session.candidate_accepted("001", 3))
        self.assertEqual(self.fixture.session().items["001"]["row"][PRICE], 11)

    def test_remove_all_candidates_clears_price_and_approval(self):
        self.session.apply_candidates("001", {"0": False, "1": False, "2": False})
        item = self.session.items["001"]
        self.assertIsNone(item["row"][PRICE])
        self.assertEqual(item["counts"], {})
        self.assertEqual(item["score"], 1)
        with self.assertRaises(ValueError):
            self.session.decide(["001"], "accepted")

    def test_validation_and_save_failure_are_atomic(self):
        before = deepcopy(self.session.items["001"])
        for overrides in ({"4": True}, {"99": True}, {"0": "yes"}):
            with self.assertRaises(ValueError):
                self.session.apply_candidates("001", overrides)
            self.assertEqual(self.session.items["001"], before)
        with patch.object(self.session, "save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                self.session.apply_candidates("001", {"3": True})
        self.assertEqual(self.session.items["001"], before)
        self.assertFalse(self.session.undo_stack)

    def test_economy_median_and_all_candidate_modes(self):
        row = self.session.original_rows["001"]
        row["Market Price Level"] = "economy"
        # Economy: Rey's 25th percentile is 12.5, Riba's is 40.
        self.assertAlmostEqual(self.session.preview_candidates("001", {})["row"][PRICE], 10.5)
        row["Average Mode"] = "all_candidates"
        self.assertAlmostEqual(self.session.preview_candidates("001", {})["row"][PRICE], 6)
        row["Market Price Level"] = "median"
        self.assertAlmostEqual(self.session.preview_candidates("001", {})["row"][PRICE], 8)

    def test_workbook_settings_override_current_app_settings(self):
        self.session.settings.update(BAP_MARKET_PRICE_PERCENT=.9, STORE_WEIGHT_CAP=99, ECONOMY_PERCENTILE=.8)
        self.assertAlmostEqual(self.session.preview_candidates("001", {})["row"][PRICE], 11)

    def test_missing_size_and_normalized_price_fallback(self):
        candidate = self.session.candidates["001"][0]
        candidate["Price / kg"] = None
        candidate["Chosen Price"] = 20
        candidate["Normalized kg"] = 2
        self.assertEqual(self.session.candidate_unit_price(candidate), 10)
        self.assertAlmostEqual(self.session.preview_candidates("001", {})["row"][PRICE], 11)
        self.session.original_rows["001"]["Target Quantity kg"] = None
        with self.assertRaisesRegex(ValueError, "target package quantity"):
            self.session.preview_candidates("001", {})

    def test_changed_workbook_does_not_reuse_candidate_indices(self):
        self.session.apply_candidates("001", {"3": True})
        self.fixture.rows[0][PRICE] = 12
        self.fixture.write_results()
        reopened = self.fixture.session()
        self.assertTrue(reopened.archived_review.exists())
        self.assertEqual(reopened.items["001"]["candidate_overrides"], {})
        self.assertEqual(reopened.items["001"]["row"][PRICE], 12)


class CandidateReviewUITests(unittest.TestCase):
    def test_filters_sort_shortcuts_and_drag_in_both_languages(self):
        fixture = candidate_fixture()
        root = tk.Tk()
        root.withdraw()
        try:
            for lang in ("es", "en"):
                review = ReviewWindow(root, fixture.results, fixture.template, lang=lang)
                review.attributes("-alpha", 0.0)
                review.deiconify()
                root.update()
                dialog = CandidateReviewWindow(review, "001")
                dialog.attributes("-alpha", 0.0)
                root.update()
                dialog.notebook.select(dialog.tabs["ribasmith"])
                root.update()
                tree = dialog.trees["ribasmith"]
                initial_preview = dialog.preview_label.cget("text")
                dialog.original_filter.set(dialog.original_values[2])
                self.assertEqual(tree.get_children(), ("3", "4"))
                dialog.search.set("Rice 3")
                self.assertEqual(tree.get_children(), ("3",))
                self.assertEqual(dialog.preview_label.cget("text"), initial_preview)
                tree.focus_force()
                root.update()
                tree.event_generate("<Control-a>")
                root.update()
                self.assertEqual(tree.selection(), ("3",))
                tree.event_generate("<KeyPress-a>")
                root.update()
                self.assertTrue(dialog.overrides["3"])
                self.assertNotIn("4", dialog.overrides)
                self.assertNotIn("0", dialog.overrides)
                self.assertIn("$13.00", dialog.preview_label.cget("text"))
                dialog.filter.set(dialog.filter_values[3])
                self.assertEqual(tree.get_children(), ("3",))
                tree.event_generate("<Control-z>")
                root.update()
                self.assertFalse(dialog.overrides)
                self.assertFalse(tree.get_children())
                dialog.clear_filters()
                dialog.sort_by("ribasmith", "price")
                self.assertEqual(tree.get_children(), ("2", "3", "4"))
                dialog.sort_by("ribasmith", "price")
                self.assertEqual(tree.get_children(), ("3", "2", "4"))
                self.assertIn("↓", tree.heading("price", "text"))
                # Drag follows the displayed sort order and can add to selection with Ctrl.
                tree.selection_set("4")
                with patch.object(tree, "identify_region", return_value="cell"), \
                     patch.object(tree, "identify_row", return_value="3"):
                    dialog.drag_start(SimpleNamespace(widget=tree, x=30, y=80, state=4))
                with patch.object(tree, "identify_row", return_value="2"), \
                     patch.object(tree, "winfo_height", return_value=400):
                    dialog.drag_move(SimpleNamespace(y=150))
                self.assertEqual(set(tree.selection()), {"2", "3", "4"})
                with patch.object(tree, "identify_row", return_value="4"), \
                     patch.object(tree, "winfo_height", return_value=400):
                    dialog.drag_move(SimpleNamespace(y=450))
                self.assertIsNotNone(dialog._drag_timer)
                dialog.tab_changed()
                self.assertIsNone(dialog._drag_timer)
                # A/R advances to the next visible candidate after filtering.
                dialog.filter.set(dialog.filter_values[1])
                self.assertEqual(tree.get_children(), ("2",))
                tree.selection_set("2")
                dialog.tree_key(SimpleNamespace(widget=tree, keysym="r", state=0))
                self.assertFalse(tree.get_children())
                dialog.undo()
                self.assertEqual(tree.get_children(), ("2",))
                dialog.clear_filters()
                tree.selection_set("2")
                dialog.tree_key(SimpleNamespace(widget=tree, keysym="space", state=0))
                self.assertFalse(dialog.session.candidate_accepted("001", 2, dialog.overrides))
                with patch("candidate_review_ui.webbrowser.open") as open_url:
                    dialog.tree_key(SimpleNamespace(widget=tree, keysym="o", state=0))
                    open_url.assert_called_once_with("https://example.com/product")
                dialog.focus_search()
                dialog.search_entry.focus_force()
                root.update()
                previous = dict(dialog.overrides)
                dialog.search_entry.event_generate("<KeyPress-a>")
                root.update()
                self.assertEqual(dialog.search.get(), "a")
                self.assertEqual(dialog.overrides, previous)
                dialog.destroy()
                self.assertEqual(review.session.items["001"]["row"][PRICE], 11)
                review.destroy()
        finally:
            root.destroy()
            fixture.tearDown()

    def test_store_tabs_preview_cancel_apply_and_parent_undo_both_languages(self):
        fixture = candidate_fixture()
        root = tk.Tk()
        root.withdraw()
        try:
            for lang in ("es", "en"):
                review = ReviewWindow(root, fixture.results, fixture.template, lang=lang)
                review.withdraw()
                review.tree.selection_set("001")
                review.show_details()
                review.open_candidates()
                dialog = review._candidate_window
                root.update_idletasks()
                self.assertEqual(set(dialog.trees), {"rey", "ribasmith", "super99"})
                self.assertEqual(dialog.trees["rey"].get_children(), ("0", "1"))
                self.assertEqual(dialog.trees["ribasmith"].get_children(), ("2", "3", "4"))
                dialog.notebook.select(dialog.tabs["ribasmith"])
                dialog.trees["ribasmith"].selection_set("3")
                dialog.choose(True)
                self.assertIn("$13.00", dialog.preview_label.cget("text"))
                dialog.reset()
                self.assertFalse(dialog.overrides)
                self.assertIn("$11.00", dialog.preview_label.cget("text"))
                dialog.choose(True)
                dialog.destroy()
                self.assertEqual(review.session.items["001"]["row"][PRICE], 11)
                dialog = CandidateReviewWindow(review, "001")
                dialog.notebook.select(dialog.tabs["ribasmith"])
                dialog.trees["ribasmith"].selection_set("3")
                dialog.choose(True)
                dialog.apply()
                self.assertEqual(review.price.get(), "13.00")
                self.assertEqual(review.tree.set("001", "proposal"), "$13.00")
                review.run(review.session.undo)
                self.assertEqual(review.price.get(), "11.00")
                review.destroy()
        finally:
            root.destroy()
            fixture.tearDown()


if __name__ == "__main__":
    unittest.main()
