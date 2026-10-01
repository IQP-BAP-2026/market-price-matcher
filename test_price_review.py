import csv
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from price_review import ReviewSession, confidence, CODE, NAME, PRICE
from current_prices import read_current_rows, read_current_prices


def evidence(**updates):
    row = {CODE: "001", NAME: "Rice", PRICE: 10, "Store Count": 3, "Accepted Count": 9,
           "Search Confidence": "HIGH", "Quality Flag": "OK", "Price CV": 0.1,
           "Store Averages / kg": "rey=$10/kg (n=3) | super99=$10/kg (n=3) | ribasmith=$10/kg (n=3)"}
    row.update(updates)
    return row


class ConfidenceTests(unittest.TestCase):
    def score(self, row, current=10, settings=None):
        return confidence(row, current, {"a": 3, "b": 3, "c": 3}, settings)

    def test_strong_evidence(self):
        self.assertEqual(self.score(evidence())[0], 5)

    def test_absolute_change_boundary_and_config(self):
        for value in (8, 12):
            self.assertEqual(self.score(evidence(**{PRICE: value}))[0], 5)
        for value in (7.99, 12.01):
            self.assertLessEqual(self.score(evidence(**{PRICE: value}))[0], 3)
        self.assertEqual(self.score(evidence(**{PRICE: 12.01}), settings={"PRICE_CHANGE_REVIEW_THRESHOLD": .3})[0], 5)

    def test_embedded_flag_details_preserve_caps(self):
        for flag in ("VERY HIGH RISK: only 2 clean match(es)",
                     "TARGET SIZE WARNING: name parses to 2kg but spreadsheet weight is 4kg"):
            self.assertLessEqual(self.score(evidence(**{"Quality Flag": flag}))[0], 2)
        self.assertLessEqual(self.score(evidence(**{"Quality Flag": "HIGH RISK: only 3 clean match(es); need 5 | PARTIAL REQUEST ERROR"}))[0], 3)

    def test_missing_and_nonfinite_prices(self):
        for price in (None, 0, -1, float("nan"), float("inf"), "abc"):
            self.assertEqual(self.score(evidence(**{PRICE: price}))[0], 1)
        self.assertLess(self.score(evidence(), current=0)[1], self.score(evidence())[1])

    def test_store_concentration_and_query_affect_score(self):
        row = evidence(**{"Store Count": 1, "Search Confidence": "LOW"})
        result = confidence(row, 10, {"a": 100})
        self.assertLess(result[0], self.score(evidence())[0])


class SessionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.template = self.root / "current.xlsx"
        self.results = self.root / "results.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Prices"
        ws.append(["ProductCode", "Name", "Id", "IsActive", "LastModifiedDate", "UnitPrice", "UseStandardPrice", "Pricebook2Id"])
        for key, name, price in (("001", "Rice", 10), ("002", "Beans", 5), ("003", "Bath mat", None), ("004", "Milk", 4)):
            ws.append([key, name, "01u1K00000aA8" + key + "AK", True, "2026-09-28T21:46:52.000Z", price, False, "01s41000004hcm1AAA"])
        ws["E2"].number_format = "$0.00"
        wb.save(self.template)
        wb.close()
        self.rows = [evidence(), evidence(**{CODE: "002", NAME: "Beans", PRICE: 6}),
                     evidence(**{CODE: "003", NAME: "Bath mat", PRICE: None})]
        self.write_results()

    def tearDown(self):
        self.temp.cleanup()

    def write_results(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        headers = list(self.rows[0])
        ws.append(headers)
        for row in self.rows:
            ws.append([row.get(h) for h in headers])
        wb.save(self.results)
        wb.close()

    def session(self):
        return ReviewSession(self.results, self.template)

    def test_default_export_only_explicit_approvals_and_exact_csv_headers(self):
        session = self.session()
        session.decide(["001"], "accepted")
        destination = self.root / "export.csv"
        self.assertEqual(session.export(destination), 1)
        with destination.open(encoding="utf-8", newline="") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows, [["Id", "UnitPrice"], ["01u1K00000aA8001AK", "10.00"]])

    def test_manual_rejected_and_missing_prices(self):
        session = self.session()
        session.decide(["001", "002"], "rejected")
        session.decide(["003"], "manual", "7.25", "Supplier invoice")
        self.assertEqual(session.export(self.root / "export.csv"), 1)
        reopened = self.session()
        self.assertEqual(reopened.items["003"]["price"], 7.25)
        self.assertEqual(reopened.items["001"]["decision"], "rejected")

    def test_bulk_atomic_validation_and_undo(self):
        session = self.session()
        with self.assertRaises(ValueError):
            session.decide(["001", "003"], "accepted")
        self.assertEqual(session.items["001"]["decision"], "pending")
        session.decide(["001", "002"], "accepted")
        session.undo()
        self.assertTrue(all(i["decision"] == "pending" for i in self.session().items.values()))

    def test_include_unchanged_retains_current_not_rejected_proposal(self):
        session = self.session()
        session.decide(["002"], "rejected")
        destination = self.root / "export.csv"
        self.assertEqual(session.export(destination, True), 3)
        with destination.open(encoding="utf-8", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual({r["Id"]: float(r["UnitPrice"]) for r in rows},
                         {"01u1K00000aA8001AK": 10, "01u1K00000aA8002AK": 5, "01u1K00000aA8004AK": 4})


    def test_manual_validation(self):
        session = self.session()
        for value, note in ((0, "reason"), (-1, "reason"), (float("nan"), "reason"), (0.001, "reason")):
            with self.assertRaises(ValueError):
                session.decide(["003"], "manual", value, note)
        session.decide(["003"], "manual", 2)
        self.assertEqual(session.items["003"]["price"], 2)

    def test_stale_reviews_are_archived_and_inputs_protected(self):
        session = self.session()
        session.decide(["001"], "accepted")
        with self.assertRaises(ValueError):
            session.export(self.template)
        self.rows[0][PRICE] = 12
        self.write_results()
        with self.assertRaises(ValueError):
            session.export(self.root / "export.csv")
        reopened = self.session()
        self.assertTrue(reopened.archived_review.is_file())
        self.assertEqual(reopened.items["001"]["decision"], "pending")

    def test_duplicate_codes_fail(self):
        self.rows.append(self.rows[0])
        self.write_results()
        with self.assertRaises(ValueError):
            self.session()

    def test_csv_review_uses_summary_per_store_counts(self):
        path = self.root / "summary.csv"
        with path.open("w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=list(self.rows[0]))
            writer.writeheader()
            writer.writerows(self.rows)
        session = ReviewSession(path, self.template)
        self.assertEqual(sum(session.items["001"]["counts"].values()), 9)
        self.assertEqual(session.items["001"]["score"], 5)

    def test_missing_template_row_blocks_without_product_metadata(self):
        self.rows.append(evidence(**{CODE: "999", NAME: "New product"}))
        self.write_results()
        session = self.session()
        session.decide(["999"], "accepted")
        with self.assertRaises(ValueError):
            session.export(self.root / "export.csv")

    def test_new_product_metadata_survives_resume_and_export(self):
        self.rows.append(evidence(**{CODE: "999", NAME: "New product"}))
        self.write_results()
        products = self.root / "products.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append([CODE, NAME, "Requisitos de Temperatura"])
        ws.append(["999", "New product", "Cold"])
        wb.save(products)
        wb.close()
        session = ReviewSession(self.results, self.template, products=products)
        session.decide(["999"], "manual", 9, "Confirmed")
        resumed = self.session()
        destination = self.root / "new.csv"
        self.assertEqual(resumed.product_rows["999"]["Requisitos de Temperatura"], "Cold")
        with self.assertRaisesRegex(ValueError, "missing Salesforce.*999"):
            resumed.export(destination)
        self.assertFalse(destination.exists())


    def test_failed_save_rolls_back_decision(self):
        session = self.session()
        with patch.object(session, "save", side_effect=PermissionError("locked")):
            with self.assertRaises(PermissionError):
                session.decide(["001"], "accepted")
        self.assertEqual(session.items["001"]["decision"], "pending")

    def test_reordered_salesforce_csv_and_all_comparison_readers(self):
        path = self.root / "salesforce.csv"
        path.write_text('UnitPrice,Id,Name,ProductCode\n200,01u1K00000aA8DoQAK,Embutidos,EMBV 17kg-CAJA\n', encoding="utf-8-sig")
        from robot_launcher import _read_current_prices
        self.assertEqual(read_current_prices(path), {"EMBV 17kg-CAJA": 200})
        self.assertEqual(_read_current_prices(path), {"EMBV 17kg-CAJA": 200})
        self.assertEqual(read_current_rows(path)[0]["Id"], "01u1K00000aA8DoQAK")

    def test_duplicate_current_codes_and_ids_are_rejected(self):
        path = self.root / "salesforce.csv"
        for body in ("a,id1,1\na,id2,2", "a,id1,1\nb,id1,2"):
            path.write_text("ProductCode,Id,UnitPrice\n" + body, encoding="utf-8")
            with self.assertRaises(ValueError):
                read_current_prices(path)

    def test_missing_id_blocks_export_but_allows_review(self):
        wb = load_workbook(self.template)
        wb.active["C2"] = None
        wb.save(self.template)
        wb.close()
        session = self.session()
        session.decide(["001"], "accepted")
        with self.assertRaisesRegex(ValueError, "missing Salesforce.*001"):
            session.export(self.root / "export.csv")
        self.assertFalse((self.root / "export.csv").exists())

    def test_export_and_reopen_work_after_current_prices_file_is_gone(self):
        session = self.session()
        session.decide(["001"], "accepted")
        moved = self.root / "moved.xlsx"
        self.template.rename(moved)
        self.assertGreater(session.export(self.root / "export.csv"), 0)
        reopened = ReviewSession(self.results, self.template)
        self.assertEqual(reopened.items["001"]["decision"], "accepted")
        self.assertTrue(reopened.base["001"]["Id"])
        self.assertGreater(reopened.export(self.root / "export2.csv"), 0)
        reopened.path.unlink()   # without a saved review there is nothing to fall back on
        with self.assertRaisesRegex(ValueError, "Salesforce Ids"):
            ReviewSession(self.results, self.template)
        moved.rename(self.template)

    def test_export_refuses_when_results_file_changes(self):
        session = self.session()
        session.decide(["001"], "accepted")
        with self.results.open("ab") as stream:
            stream.write(b" ")
        with self.assertRaisesRegex(ValueError, "results file changed"):
            session.export(self.root / "export.csv")

    def test_redo_after_undo_and_new_change_clears_redo(self):
        session = self.session()
        session.decide(["001"], "accepted")
        session.decide(["002"], "rejected")
        session.undo()
        session.undo()
        self.assertEqual([session.items[k]["decision"] for k in ("001", "002")], ["pending", "pending"])
        session.redo()
        self.assertEqual(session.items["001"]["decision"], "accepted")
        session.redo()
        self.assertEqual(session.items["002"]["decision"], "rejected")
        session.redo()   # nothing left: no change
        session.undo()
        self.assertEqual(session.items["002"]["decision"], "pending")
        session.decide(["003"], "rejected")
        self.assertFalse(session.redo_stack)
        reopened = self.session()
        self.assertEqual(reopened.items["001"]["decision"], "accepted")

    def test_accept_without_proposal_keeps_current_bap_price(self):
        self.rows.append(evidence(**{CODE: "004", NAME: "Milk", PRICE: None}))
        self.write_results()
        session = self.session()
        session.decide(["004"], "accepted")
        self.assertEqual((session.items["004"]["decision"], session.items["004"]["price"]), ("accepted", 4))
        session.decide(["001", "004"], "accepted")   # also in a group
        self.assertEqual(session.items["004"]["price"], 4)
        with self.assertRaisesRegex(ValueError, "no proposed or current price"):
            session.decide(["003"], "accepted")      # neither a proposal nor a current price
        session.export(self.root / "export.csv")
        with (self.root / "export.csv").open(encoding="utf-8") as stream:
            rows = list(csv.reader(stream))
        self.assertIn(["01u1K00000aA8004AK", "4.00"], rows)

    def test_results_with_salesforce_ids_review_and_export_without_current_prices(self):
        currents = {"001": 10, "002": 5, "003": None}
        for row in self.rows:
            row["Current BAP Price"] = currents[row[CODE]]
            row["Salesforce Id"] = "01u1K00000aA8" + row[CODE] + "AK"
        self.write_results()
        session = ReviewSession(self.results)              # no current-prices file at all
        self.assertEqual(session.items["001"]["current"], 10)
        session.decide(["001"], "accepted")
        session.decide(["002"], "manual", "5.50")
        self.assertEqual(session.export(self.root / "export.csv"), 2)
        with (self.root / "export.csv").open(encoding="utf-8") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[0], ["Id", "UnitPrice"])
        self.assertIn(["01u1K00000aA8002AK", "5.50"], rows)
        # opening the same results later with a current-prices file keeps the decisions
        reopened = ReviewSession(self.results, self.template)
        self.assertIsNone(reopened.archived_review)
        self.assertEqual(reopened.items["002"]["decision"], "manual")

    def test_launcher_writes_salesforce_id_into_results(self):
        from openpyxl import Workbook
        from robot_launcher import _add_columns_to_workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws.append(["Spreadsheet Row", "Código de producto"])
        ws.append([2, "001"])
        ws.append([3, "002"])
        _add_columns_to_workbook(wb, {"2": (10, 0.1, "", "01uID001")}, 0.2)
        headers = [c.value for c in ws[1]]
        self.assertIn("Salesforce Id", headers)
        col = headers.index("Salesforce Id") + 1
        self.assertEqual((ws.cell(2, col).value, ws.cell(3, col).value), ("01uID001", None))

    def test_legacy_comparison_still_works(self):
        path = self.root / "legacy.csv"
        path.write_text("Código de producto,Nombre del producto,Precio de lista\n001,Rice,5\n", encoding="utf-8")
        self.assertEqual(read_current_prices(path), {"001": 5})

    def test_brief_notes_hide_raw_diagnostics(self):
        from price_review import reviewer_notes
        item = self.session().items["001"]
        item["row"]["Request Error"] = "HTTP traceback secret technical noise"
        notes = reviewer_notes(item, {}, {"Id": "entry"})
        self.assertTrue(any("failed" in n for n in notes))
        self.assertNotIn("traceback", " ".join(notes))
        self.assertLess(len(notes), 10)


if __name__ == "__main__":
    unittest.main()
