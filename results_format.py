"""How the results and performance-analysis workbooks look: Spanish sheet names and headers, which
columns BAP sees, and the reverse mapping so the app can read the files back.

The robot and the review work with the internal (English) column keys everywhere. Only the files
use the Spanish names below. Columns the review needs but BAP doesn't (search terms, confidence
inputs, calculation settings, Salesforce Id…) are kept at the right of the sheet, hidden. Files
written by older versions (English headers) still read the same way.
"""
from __future__ import annotations

import re

# ---------------------------------------------------------------------------
# Sheet names
# ---------------------------------------------------------------------------
SUMMARY_SHEET = "Resumen"
CANDIDATES_SHEET = "Candidatos"
DASHBOARD_SHEET = "Panel"
ANALYSIS_SHEET = "Análisis de productos"
LEGACY_SHEETS = {SUMMARY_SHEET: "Summary", CANDIDATES_SHEET: "Candidates",
                 DASHBOARD_SHEET: "Dashboard", ANALYSIS_SHEET: "Product Analysis"}


def find_sheet(wb, name):
    """The sheet called `name` (or its English name in files from older versions), or None."""
    for candidate in (name, LEGACY_SHEETS.get(name)):
        if candidate and candidate in wb.sheetnames:
            return wb[candidate]
    return None


# ---------------------------------------------------------------------------
# Column names: internal key -> Spanish header
# ---------------------------------------------------------------------------
SUMMARY_ES = {
    "Spreadsheet Row": "Fila en la hoja de productos",
    "ERP Id (QBO)": "ERP Id (QBO)",
    "Código de producto": "Código de producto",
    "Nombre del producto": "Nombre del producto",
    "Generated Queries": "Búsquedas",
    "Search Confidence": "Confianza de la búsqueda",
    "Search Warning": "Aviso de la búsqueda",
    "Context Modifiers": "Modificadores de contexto",
    "Selected Stores": "Supermercados consultados",
    "Stores Used": "Supermercados usados (código)",
    "Store Count": "N.º de supermercados",
    "Average Mode": "Modo de promedio",
    "Store Averages / kg": "Promedio por supermercado (por kg)",
    "Raw Identity Matches": "Coincidencias iniciales",
    "Size-Distant Removed": "Descartados por tamaño",
    "Duplicates Removed": "Duplicados descartados",
    "Price Outliers Removed": "Precios atípicos descartados",
    "Accepted Count": "Productos comparables",
    "Rejected Count": "Productos descartados",
    "Target Quantity kg": "Peso del producto (kg)",
    "Target Quantity Source": "Origen del peso",
    "Average Price / kg": "Precio promedio por kg",
    "Median Price / kg": "Precio mediano por kg",
    "Economy Price / kg": "Precio económico por kg",
    "Market Price Level": "Nivel de precio de mercado",
    "Estimated New Product Price": "Precio BAP propuesto",
    "Min Price / kg": "Precio mínimo por kg",
    "Max Price / kg": "Precio máximo por kg",
    "Price Ratio": "Relación máximo / mínimo",
    "Price CV": "Variación de precios (CV)",
    "Price Basis": "Base de precio",
    "Quality Flag": "Alertas (código)",
    "Request Error": "Error de supermercado",
    "Notes": "Notas técnicas",
    "BAP_MARKET_PRICE_PERCENT": "Porcentaje BAP",
    "STORE_WEIGHT_CAP": "Límite de peso por supermercado",
    "ECONOMY_PERCENTILE": "Percentil económico",
    # added when the search compares with the current prices
    "Current BAP Price": "Precio BAP actual",
    "% Difference vs Current": "Diferencia vs. actual",
    "Large Price Change": "Cambio grande",          # the header also shows the limit, see change_header()
    "Salesforce Id": "Id de Salesforce",
    # performance analysis only
    "Needs Manual Review": "Necesita revisión",
    "Review Reason": "Motivo de revisión",
}

DETAIL_ES = {
    "Spreadsheet Row": "Fila en la hoja de productos",
    "Código de producto": "Código de producto",
    "Target Product": "Producto BAP",
    "Store": "Supermercado (código)",
    "Store Name": "Supermercado",
    "Search Query": "Búsqueda",
    "Status": "Estado",
    "SKU": "SKU",
    "Candidate Product": "Producto en el supermercado",
    "Product URL": "Enlace",
    "Regular Price": "Precio regular",
    "Final Price": "Precio final",
    "Chosen Price": "Precio",
    "Price Basis": "Base de precio",
    "Discount %": "Descuento",
    "Currency": "Moneda",
    "Normalized kg": "Tamaño (kg)",
    "Target Unit kg": "Tamaño por unidad BAP (kg)",
    "Candidate Unit kg": "Tamaño por unidad en tienda (kg)",
    "Size Ratio": "Relación de tamaño",
    "Price / kg": "Precio por kg",
    "Quantity Parsed From": "Tamaño leído de",
    "Reason": "Motivo (técnico)",
}

# Columns computed only for people to read (never read back by the app).
DISPLAY_ES = {
    "=market_kg": "Precio de mercado por kg",
    "=stores": "Supermercados",
    "=alerts": "Alertas",
    "=reason": "Motivo",
}

# Values written in Spanish, and read back to the internal value.
VALUE_ES = {
    "Status": {"ACCEPTED": "Aceptado", "REJECTED": "Rechazado"},
    "Large Price Change": {"YES": "SÍ"},
    "Needs Manual Review": {"YES": "SÍ", "NO": "NO"},
}

STORE_NAMES = {"super99": "Super 99", "superxtra": "Super Xtra", "rey": "El Rey", "ribasmith": "Riba Smith"}

MONEY, PERCENT, CHANGE, KG = "$0.00", "0.0%", "+0.0%;-0.0%;0.0%", "0.000"
FORMATS = {
    "Estimated New Product Price": MONEY, "Current BAP Price": MONEY, "=market_kg": MONEY,
    "Average Price / kg": MONEY, "Median Price / kg": MONEY, "Economy Price / kg": MONEY,
    "Min Price / kg": MONEY, "Max Price / kg": MONEY, "Price CV": PERCENT, "% Difference vs Current": CHANGE,
    "Target Quantity kg": KG, "Chosen Price": MONEY, "Price / kg": MONEY, "Regular Price": MONEY,
    "Final Price": MONEY, "Normalized kg": KG, "Target Unit kg": KG, "Candidate Unit kg": KG,
    "Discount %": PERCENT, "Size Ratio": "0.00x",
}

# ---------------------------------------------------------------------------
# What the results file shows. (key, visible) in sheet order; keys that aren't in the data are skipped.
# Hidden columns are kept because the review reads them; everything else is left out of this file
# (the performance analysis keeps every column).
# ---------------------------------------------------------------------------
SUMMARY_LAYOUT = [
    ("Código de producto", True), ("Nombre del producto", True), ("Target Quantity kg", True),
    ("Current BAP Price", True), ("Estimated New Product Price", True), ("% Difference vs Current", True),
    ("Large Price Change", True), ("=market_kg", True), ("=stores", True), ("Accepted Count", True),
    ("=alerts", True),
    # read by the review
    ("Spreadsheet Row", False), ("Salesforce Id", False), ("Generated Queries", False),
    ("Search Confidence", False), ("Search Warning", False), ("Selected Stores", False), ("Stores Used", False),
    ("Store Count", False), ("Average Mode", False), ("Store Averages / kg", False), ("Rejected Count", False),
    ("Target Quantity Source", False), ("Average Price / kg", False), ("Median Price / kg", False),
    ("Economy Price / kg", False), ("Min Price / kg", False), ("Max Price / kg", False),
    ("Market Price Level", False), ("Price Ratio", False), ("Price CV", False), ("Quality Flag", False),
    ("Request Error", False), ("Notes", False), ("BAP_MARKET_PRICE_PERCENT", False),
    ("STORE_WEIGHT_CAP", False), ("ECONOMY_PERCENTILE", False),
]

CANDIDATES_LAYOUT = [
    ("Código de producto", True), ("Target Product", True), ("Store Name", True), ("Candidate Product", True),
    ("Status", True), ("Chosen Price", True), ("Normalized kg", True), ("Price / kg", True), ("=reason", True),
    ("Product URL", True),
    # read by the review's "Choose candidates" window
    ("Spreadsheet Row", False), ("Store", False), ("Search Query", False), ("SKU", False), ("Reason", False),
]

# The limit shown in the "big change" header; robot_launcher sets it from the settings.
CHANGE_LIMIT = 0.20


def change_header(limit=None):
    return f"Cambio mayor a ±{(CHANGE_LIMIT if limit is None else limit):.0%}"


def header_es(key, kind="summary"):
    if key == "Large Price Change":
        return change_header()
    if key in DISPLAY_ES:
        return DISPLAY_ES[key]
    table = DETAIL_ES if kind == "detail" else SUMMARY_ES
    return table.get(key, key)


def value_es(key, value):
    return VALUE_ES.get(key, {}).get(value, value) if isinstance(value, str) else value


# ---------------------------------------------------------------------------
# Plain-language text for the visible columns
# ---------------------------------------------------------------------------
def store_names(ids):
    return ", ".join(STORE_NAMES.get(s.strip(), s.strip()) for s in str(ids or "").split(",") if s.strip())


_ALERTS = [
    (r"^OK$", lambda m: ""),
    (r"^NO MATCHES", lambda m: "Sin productos comparables"),
    (r"^VERY HIGH RISK: only (\d+)", lambda m: f"Riesgo muy alto: solo {m[1]} producto(s) comparable(s)"),
    (r"^HIGH RISK: only (\d+).*?need (\d+)", lambda m: f"Riesgo alto: solo {m[1]} productos comparables (se buscan {m[2]})"),
    (r"^HIGH RISK: only (\d+)", lambda m: f"Riesgo alto: solo {m[1]} productos comparables"),
    (r"^ONE STORE ONLY", lambda m: "Un solo supermercado"),
    (r"^LOW SEARCH CONFIDENCE", lambda m: "Búsqueda poco precisa"),
    (r"^TARGET SIZE WARNING: name parses to ([\d.]+)kg but spreadsheet weight is ([\d.]+)kg",
     lambda m: f"Revisar el peso: el nombre indica {float(m[1]):g} kg y la hoja {float(m[2]):g} kg"),
    (r"^TARGET SIZE WARNING", lambda m: "Revisar el peso del producto"),
    (r"^PARTIAL REQUEST ERROR", lambda m: "Falló la consulta a algún supermercado"),
    (r"^REQUEST ERROR", lambda m: "Falló la consulta a los supermercados"),
]
_ALERTS = [(re.compile(pattern), build) for pattern, build in _ALERTS]


def alerts_es(quality_flag):
    """The robot's quality flags ("HIGH RISK: only 3 clean match(es); need 5 | ONE STORE ONLY") in Spanish."""
    parts = []
    for flag in str(quality_flag or "").split("|"):
        flag = flag.strip()
        if not flag:
            continue
        for pattern, build in _ALERTS:
            match = pattern.search(flag)
            if match:
                flag = build(match)
                break
        if flag:
            parts.append(flag)
    return " · ".join(parts)


def _number(value):
    try:
        return float(value) if value not in (None, "") and not isinstance(value, bool) else None
    except (TypeError, ValueError):
        return None


def _market_kg(row):
    """The market price per kg behind the proposal (proposal = market price/kg × kg × BAP %)."""
    price, kg, percent = (_number(row.get(k)) for k in ("Estimated New Product Price", "Target Quantity kg",
                                                        "BAP_MARKET_PRICE_PERCENT"))
    return price / (kg * percent) if price is not None and kg and percent else None


DISPLAY_VALUES = {
    "=market_kg": _market_kg,
    "=stores": lambda row: store_names(row.get("Stores Used")),
    "=alerts": lambda row: alerts_es(row.get("Quality Flag")),
    "=reason": lambda row: friendly_reason(row.get("Reason"), "es"),
}


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------
class Table:
    """`rows` (lists in the order of the internal `headers`) laid out as `layout` says, with Spanish
    headers and values. Rows are converted one at a time, so even the large Candidates sheet isn't
    copied in memory."""

    def __init__(self, headers, rows, layout, kind="summary"):
        index = {h: i for i, h in enumerate(headers)}
        self.columns = [(key, visible) for key, visible in layout if key in index or key in DISPLAY_VALUES]
        self.headers = [header_es(key, kind) for key, _ in self.columns]
        self.hidden = [not visible for _, visible in self.columns]
        self.formats = [FORMATS.get(key, "General") for key, _ in self.columns]
        self._index, self._source, self._rows = index, list(headers), rows

    def convert(self, row):
        record = None
        out = []
        for key, _visible in self.columns:
            if key in DISPLAY_VALUES:
                if record is None:
                    record = dict(zip(self._source, row))
                out.append(DISPLAY_VALUES[key](record))
            else:
                i = self._index[key]
                out.append(value_es(key, row[i] if i < len(row) else None))
        return out

    def __len__(self):
        return len(self._rows)

    def __iter__(self):
        return (self.convert(row) for row in self._rows)

    def __getitem__(self, item):
        if isinstance(item, slice):
            return [self.convert(row) for row in self._rows[item]]
        return self.convert(self._rows[item])


def full_table(headers, rows, kind="summary", extra_display=()):
    """Every column, translated (for the performance analysis)."""
    layout = [(h, True) for h in headers] + [(key, True) for key in extra_display]
    return Table(headers, rows, layout, kind)


def write_sheet(ws, table, style_sheet):
    """Write a (small) table into an openpyxl sheet: Spanish headers, number formats, hidden columns."""
    from openpyxl.utils import get_column_letter
    ws.append(table.headers)
    for row in table:
        ws.append(row)
    style_sheet(ws)
    finish_columns(ws, table)
    for col, fmt in enumerate(table.formats, 1):
        if fmt != "General":
            for r in range(2, ws.max_row + 1):
                ws.cell(r, col).number_format = fmt
    # big price changes stand out, as in the app
    keys = [key for key, _ in table.columns]
    if "Large Price Change" in keys:
        from openpyxl.styles import Font, PatternFill
        fill, font = PatternFill("solid", fgColor="FDE2C4"), Font(bold=True, color="9A3412")
        flag = keys.index("Large Price Change") + 1
        diff = keys.index("% Difference vs Current") + 1 if "% Difference vs Current" in keys else None
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, flag).value == value_es("Large Price Change", "YES"):
                ws.cell(r, flag).fill, ws.cell(r, flag).font = fill, font
                if diff:
                    ws.cell(r, diff).fill = fill
    return {get_column_letter(i + 1): h for i, h in enumerate(table.headers)}


def finish_columns(ws, table):
    """Hide the columns only the app reads; keep the filter on the visible ones."""
    from openpyxl.utils import get_column_letter
    visible = sum(1 for hidden in table.hidden if not hidden)
    for col, hidden in enumerate(table.hidden, 1):
        if hidden:
            ws.column_dimensions[get_column_letter(col)].hidden = True
    if visible and visible < len(table.hidden):
        ws.auto_filter.ref = f"A1:{get_column_letter(visible)}{max(ws.max_row, 2)}"
    # printing: landscape, all visible columns on one page width, header row on every page
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


# ---------------------------------------------------------------------------
# Reading (results files from this version and from older ones)
# ---------------------------------------------------------------------------
_REVERSE = {}
for _table in (SUMMARY_ES, DETAIL_ES):
    for _key, _es in _table.items():
        _REVERSE.setdefault(_es, _key)
for _key, _es in DISPLAY_ES.items():
    _REVERSE[_es] = _key
_VALUE_REVERSE = {key: {es: value for value, es in values.items()} for key, values in VALUE_ES.items()}


def internal_header(header):
    text = "" if header is None else str(header)
    if text in _REVERSE:
        return _REVERSE[text]
    if text.startswith("Cambio mayor a") or text.startswith("Price Change >"):
        return "Large Price Change"
    return text


def internal_record(headers, row):
    """One row of a results file as {internal key: internal value}."""
    keys = [internal_header(h) for h in headers]
    record = {}
    for key, value in zip(keys, row):
        if key in _VALUE_REVERSE and isinstance(value, str):
            value = _VALUE_REVERSE[key].get(value, value)
        record[key] = value
    return record


def read_records(path, sheet_name=SUMMARY_SHEET):
    """All data rows of a results file (.xlsx or .csv) as internal records."""
    from pathlib import Path
    path = Path(path)
    from current_prices import is_spreadsheet, read_text_table
    if not is_spreadsheet(path):
        rows, _ = read_text_table(path)
        headers, body = (rows[0], rows[1:]) if rows else ([], [])
        return headers, [internal_record(headers, r) for r in body if any(v not in (None, "") for v in r)]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = find_sheet(wb, sheet_name) or wb.active
        rows = ws.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        return headers, [internal_record(headers, r) for r in rows if any(v is not None for v in r)]
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Plain-language reasons. The robot writes short technical notes ("SIZE-DISTANT MATCH REMOVED: ratio=0.42; …");
# each rule turns one kind of note into a sentence. Anything not recognised is shown as written.
# ---------------------------------------------------------------------------
def _pct(value):
    try:
        return f"{float(value) * 100:.0f} %"
    except (TypeError, ValueError):
        return "?"


_REASON_RULES = [
    (r"^accepted; sold by weight", lambda m: ("Aceptado: se vende por peso", "Accepted: sold by weight")),
    (r"^accepted; size strategy=PREFERRED.*?ratio=([\d.]+)",
     lambda m: (f"Aceptado: tamaño parecido ({_pct(m[1])} del producto BAP)",
                f"Accepted: similar size ({_pct(m[1])} of the BAP product)")),
    (r"^accepted; size strategy=CLOSE.*?ratio=([\d.]+)",
     lambda m: (f"Aceptado: tamaño cercano ({_pct(m[1])} del producto BAP)",
                f"Accepted: close size ({_pct(m[1])} of the BAP product)")),
    (r"^accepted; size strategy=FALLBACK.*?ratio=([\d.]+)",
     lambda m: (f"Aceptado: el tamaño más cercano disponible ({_pct(m[1])} del producto BAP)",
                f"Accepted: nearest size available ({_pct(m[1])} of the BAP product)")),
    (r"^accepted", lambda m: ("Aceptado", "Accepted")),
    (r"SIZE-DISTANT MATCH REMOVED: ratio=([\d.]+)",
     lambda m: (f"Tamaño muy distinto ({_pct(m[1])} del producto BAP)",
                f"Size too different ({_pct(m[1])} of the BAP product)")),
    (r"SIZE PEER REMOVED", lambda m: ("Sin tamaño, mientras otros productos sí lo tienen",
                                      "No size, while other listings have one")),
    (r"cannot normalize candidate to kg", lambda m: ("No se pudo leer el tamaño del paquete",
                                                      "The package size couldn't be read")),
    (r"cannot count units", lambda m: ("No se pudieron contar las unidades del paquete",
                                       "The number of units couldn't be read")),
    (r"missing usable retailer price", lambda m: ("El supermercado no muestra un precio válido",
                                                   "The store shows no usable price")),
    (r"missing required identity word\(s\): (.+)",
     lambda m: (f"Le faltan palabras clave: {m[1]}", f"Missing key words: {m[1]}")),
    (r"unexpected subtype for plain '(.+?)': (.+)",
     lambda m: (f"Variante distinta de «{m[1]}»: {m[2]}", f"Different kind of “{m[1]}”: {m[2]}")),
    (r"PRICE OUTLIER REMOVED: \$([\d.,]+)/kg vs median \$([\d.,]+)/kg",
     lambda m: (f"Precio fuera de rango: ${m[1]}/kg frente a la mediana ${m[2]}/kg",
                f"Price out of range: ${m[1]}/kg vs median ${m[2]}/kg")),
    (r"different item size: target ([\d.]+), candidate ([\d.]+)",
     lambda m: (f"Tamaño por unidad distinto: BAP {m[1]}, tienda {m[2]} (kg/L)",
                f"Different unit size: BAP {m[1]}, store {m[2]} (kg/L)")),
    (r"different dimensions: target (\S+), candidate (\S+)",
     lambda m: (f"Medidas distintas: BAP {m[1]}, tienda {m[2]}", f"Different dimensions: BAP {m[1]}, store {m[2]}")),
    (r"target word only appears after '(.+?)'",
     lambda m: (f"La palabra clave solo aparece después de «{m[1]}» (ingrediente o sabor, no el producto)",
                f"The key word only appears after “{m[1]}” (ingredient or flavour, not the product)")),
    (r"different product: title leads with '(.+?)'",
     lambda m: (f"Otro producto: el nombre empieza con «{m[1]}»", f"Different product: the name starts with “{m[1]}”")),
    (r"speciality/premium version of the product: (.+)",
     lambda m: (f"Versión especial o premium: {m[1]}", f"Special or premium version: {m[1]}")),
    (r"preparation/mix, not the product itself: (.+)",
     lambda m: (f"Mezcla o preparado, no el producto: {m[1]}", f"Mix or preparation, not the product: {m[1]}")),
    (r"made from/with '(.+?)', not the product itself",
     lambda m: (f"Hecho con «{m[1]}», no es el producto", f"Made with “{m[1]}”, not the product itself")),
    (r"false[- ]positive (?:family )?term\(s\): (.+)",
     lambda m: (f"Coincidencia engañosa: {m[1]}", f"Misleading match: {m[1]}")),
    (r"alcoholic product: (.+)", lambda m: (f"Producto alcohólico: {m[1]}", f"Alcoholic product: {m[1]}")),
    (r"pet product: (.+)", lambda m: (f"Producto para mascotas: {m[1]}", f"Pet product: {m[1]}")),
    (r"non-food product for a food target: (.+)", lambda m: (f"No es un alimento: {m[1]}", f"Not a food product: {m[1]}")),
    (r"breaded/prepared version of the product: (.+)",
     lambda m: (f"Versión empanizada o preparada: {m[1]}", f"Breaded or prepared version: {m[1]}")),
    (r"processed/canned product for a fresh-produce target: (.+)",
     lambda m: (f"Procesado o enlatado (se busca fresco): {m[1]}", f"Processed or canned (fresh wanted): {m[1]}")),
    (r"cooked/deli/seasoned version of a raw meat product: (.+)",
     lambda m: (f"Versión cocida o de charcutería: {m[1]}", f"Cooked or deli version: {m[1]}")),
    (r"physical-form mismatch: (.+)",
     lambda m: ("Forma distinta (sólido frente a líquido)", "Different form (solid vs liquid)")),
    (r"different/extra flavour\(s\): (.+)", lambda m: (f"Sabor distinto o adicional: {m[1]}", f"Different or extra flavour: {m[1]}")),
    (r"NEAR-DUPLICATE REMOVED: overlaps SKU \S+ \((.+)\)",
     lambda m: (f"Duplicado de «{m[1]}»", f"Duplicate of “{m[1]}”")),
    (r"generic product word '(.+?)' appears too late",
     lambda m: (f"«{m[1]}» aparece muy tarde en el nombre", f"“{m[1]}” appears too late in the name")),
]
_REASON_RULES = [(re.compile(pattern, re.I), build) for pattern, build in _REASON_RULES]


def friendly_reason(reason, lang="es", short=False):
    """The robot's technical note as a sentence in the reviewer's language. short=True drops what the
    table already shows elsewhere ("Accepted:" — that's the Robot column — and "of the BAP product")."""
    text = str(reason or "").strip()
    for pattern, build in _REASON_RULES:
        match = pattern.search(text)
        if match:
            es, en = build(match)
            text = es if lang == "es" else en
            if short:
                text = re.sub(r"^(Aceptado|Accepted)(: |$)", "", text)
                text = text.replace(" del producto BAP", "").replace(" of the BAP product", "")
                text = text[:1].upper() + text[1:] if text else ("Aceptado" if lang == "es" else "Accepted")
            return text
    return text
