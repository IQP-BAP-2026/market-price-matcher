"""Read the files people give the robot — the products list and the current prices — whatever their format.

* .xlsx / .xlsm workbooks (read with openpyxl);
* old .xls workbooks (read with xlrd);
* "Excel" files that are really an HTML table (Salesforce report exports often are), whatever their name;
* .csv / .txt text tables, in UTF-8, UTF-16 ("Unicode text") or the Windows code page (cp1252), separated
  by commas, semicolons or tabs, and with decimal commas (5,02) when the separator is a semicolon.

The format is decided by the file's content, not its name, so a CSV saved as "precios.xlsx" still works.
Every reader returns the same small workbook interface that the robot already uses with openpyxl:
`sheetnames`, `book[name]`, `active`, `close()`, and per sheet `title`, `max_row`, `max_column`,
`cell(row, col).value` and `iter_rows(min_row=, max_row=, values_only=True)`.
"""
from __future__ import annotations

import csv
import io
from html.parser import HTMLParser
from pathlib import Path

_TEXT_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")
_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def file_kind(path):
    """'xlsx', 'xls', 'html' or 'text' — from the file's first bytes."""
    with open(path, "rb") as stream:
        start = stream.read(4096)
    if start.startswith(b"PK"):
        return "xlsx"
    if start.startswith(_OLE):
        return "xls"
    head = _decode(start).lstrip("﻿ \r\n\t").lower()
    if head.startswith(("<html", "<!doctype", "<table", "<?xml")) or "<table" in head[:2000]:
        return "html"
    return "text"


def is_spreadsheet(path):
    """True when openpyxl can open the file (an .xlsx/.xlsm workbook, whatever it's called)."""
    return file_kind(path) == "xlsx"


def _decode(raw):
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    for encoding in _TEXT_ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def read_text_table(path):
    """(rows, decimal_comma) of a CSV/text file in any of the encodings and separators Excel uses."""
    text = _decode(Path(path).read_bytes())
    lines = text.splitlines()
    if lines and lines[0].strip().lower().startswith("sep="):      # Excel's "sep=;" hint line
        delimiter = lines[0].strip()[4:5] or ","
        text = "\n".join(lines[1:])
        dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect, delimiter=delimiter))
        return rows, delimiter == ";"
    try:
        dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text, newline=""), dialect))
    return rows, dialect.delimiter == ";"


class _HTMLTable(HTMLParser):
    """The rows of the tables in an HTML file (cell text only)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows, self._row, self._cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append("\n")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._cell is not None:
                self.handle_endtag("td")
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def read_html_table(path):
    import re
    raw = Path(path).read_bytes()
    text = _decode(raw)
    declared = re.search(r"""(?:charset|encoding)\s*=\s*["']?([\w-]+)""", text[:3000], re.I)
    if declared and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):   # a declared encoding wins
        try:
            text = raw.decode(declared.group(1))
        except (LookupError, UnicodeDecodeError):
            pass
    parser = _HTMLTable()
    parser.feed(text)
    parser.close()
    return parser.rows


# ---------------------------------------------------------------------------
# A minimal in-memory workbook with the openpyxl methods the robot uses
# ---------------------------------------------------------------------------
class _Cell:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class TableSheet:
    def __init__(self, title, rows, decimal_comma=False):
        rows = [[None if (isinstance(v, str) and not v.strip()) else v for v in row] for row in rows]
        while rows and not any(v is not None for v in rows[-1]):     # trailing empty lines
            rows.pop()
        self.title = title
        self.max_column = max((len(r) for r in rows), default=0)
        self.rows = [list(r) + [None] * (self.max_column - len(r)) for r in rows]
        self.max_row = len(self.rows)
        self.decimal_comma = decimal_comma

    def cell(self, row, column):
        if 1 <= row <= self.max_row and 1 <= column <= self.max_column:
            return _Cell(self.rows[row - 1][column - 1])
        return _Cell(None)

    def iter_rows(self, min_row=1, max_row=None, values_only=True):
        last = self.max_row if max_row is None else min(max_row, self.max_row)
        for index in range(max(1, min_row) - 1, last):
            yield tuple(self.rows[index])


class TableBook:
    def __init__(self, sheets):
        self._sheets = {sheet.title: sheet for sheet in sheets}
        self.sheetnames = [sheet.title for sheet in sheets]

    def __getitem__(self, name):
        return self._sheets[name]

    def __contains__(self, name):
        return name in self._sheets

    @property
    def active(self):
        return self._sheets[self.sheetnames[0]]

    def close(self):
        pass


def _read_xls(path):
    try:
        import xlrd
    except ImportError:
        raise ValueError(f"{Path(path).name} is an old Excel (.xls) file and this copy of the app can't read it. "
                         "Open it in Excel and save it as .xlsx (Excel workbook) or .csv, then choose it again.")
    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        sheets = []
        for name in book.sheet_names():
            sheet = book.sheet_by_name(name)
            rows = []
            for r in range(sheet.nrows):
                values = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    kind, value = cell.ctype, cell.value
                    if kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        value = None
                    elif kind == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                        value = int(value)
                    elif kind == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                        except Exception:
                            pass
                    elif kind == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    values.append(value)
                rows.append(values)
            sheets.append(TableSheet(name, rows))
        return TableBook(sheets)
    finally:
        book.release_resources()


def load_book(path, read_only=True):
    """Open a products / current-prices file of any supported format (see the module notes)."""
    path = Path(path)
    kind = file_kind(path)
    if kind == "xlsx":
        from openpyxl import load_workbook
        return load_workbook(path, read_only=read_only, data_only=True)
    if kind == "xls":
        return _read_xls(path)
    if kind == "html":
        return TableBook([TableSheet(path.stem[:31] or "Hoja1", read_html_table(path))])
    rows, decimal_comma = read_text_table(path)
    return TableBook([TableSheet(path.stem[:31] or "Hoja1", rows, decimal_comma)])


def to_number(value, decimal_comma=False):
    """A number from a cell: 5.02, "$5.02", "1,234.50" — or, with decimal_comma (Spanish-format files),
    "5,02" and "1.234,50". None when it isn't a number."""
    import math
    try:
        if value is None or isinstance(value, bool):
            return None
        if isinstance(value, (int, float)):
            result = float(value)
        else:
            text = str(value).replace("$", "").replace(" ", "").replace(" ", "").strip()
            if "," in text and "." in text:      # both: the last one is the decimal separator
                comma_decimal = text.rfind(",") > text.rfind(".")
            else:                                  # only a comma: decimal in Spanish-format files
                comma_decimal = decimal_comma and "," in text
            text = text.replace(".", "").replace(",", ".") if comma_decimal else text.replace(",", "")
            result = float(text)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None
