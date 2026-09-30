from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from excel_export import prepare_candidates, save_with_candidates
from candidate_store import CandidateStore, build_candidate_index, file_hash
from price_robot import style_sheet


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "results.xlsx"
        self.headers = ["Código de producto", "Candidate Product", "Chosen Price", "Normalized kg", "Store", "Status", "SKU"]
        self.rows = [["001", ' arroz & <leche> "sí"\n ', 2.55, .25, "rey", "ACCEPTED", "0007"],
                     ["002", "=literal", None, None, "ribasmith", "REJECTED", "0099"]]

    def workbook(self, rows):
        wb = Workbook()
        wb.active.title = "Summary"
        wb.active.append(["Product", "Price"])
        wb.active.append(["001", 1.23])
        wb.active["B2"].number_format = "$0.00"
        prepare_candidates(wb.create_sheet("Candidates"), self.headers, rows, style_sheet)
        return wb

    def test_streamed_values_styles_dimensions_and_all_rows(self):
        rows = self.rows * 100
        wb = self.workbook(rows)
        self.assertLessEqual(wb["Candidates"].max_row, 150)
        save_with_candidates(wb, self.path, rows)
        restored = load_workbook(self.path)
        try:
            ws = restored["Candidates"]
            self.assertEqual(list(ws.values), [tuple(self.headers)] + [tuple(row) for row in rows])
            self.assertEqual(ws.dimensions, "A1:G201")
            self.assertEqual(ws.auto_filter.ref, "A1:G201")
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertEqual(ws["C201"].number_format, "General")  # omitted empty cell
            self.assertEqual(ws["C200"].number_format, "$0.00")
            self.assertEqual(ws["D200"].number_format, "0.000")
            self.assertEqual(ws["B3"].data_type, "s")
            self.assertTrue(ws["A1"].font.bold)
            self.assertEqual(restored["Summary"]["B2"].number_format, "$0.00")
        finally:
            restored.close()

    def test_empty_and_failed_save_preserve_previous_file(self):
        save_with_candidates(self.workbook([]), self.path, [])
        wb = load_workbook(self.path, read_only=True)
        self.assertEqual(list(wb["Candidates"].values), [tuple(self.headers)])
        wb.close()
        previous = self.path.read_bytes()
        with self.assertRaises(TypeError):
            save_with_candidates(self.workbook([]), self.path, [[object()]])
        self.assertEqual(self.path.read_bytes(), previous)
        self.assertFalse(list(self.path.parent.glob("*.tmp")))

    def test_candidate_index_is_lazy_reusable_and_invalidated(self):
        save_with_candidates(self.workbook(self.rows), self.path, self.rows)
        build_candidate_index(self.path, self.headers, self.rows)
        fingerprint = file_hash(self.path)
        with patch("openpyxl.load_workbook", side_effect=AssertionError("Should use existing index")):
            store = CandidateStore(self.path, fingerprint)
            self.assertEqual(store.counts, {"001": {"rey": 1}, "002": {}})
            self.assertEqual(store["001"][0]["SKU"], "0007")
        self.rows[0][2] = 9
        save_with_candidates(self.workbook(self.rows), self.path, self.rows)
        rebuilt = CandidateStore(self.path, file_hash(self.path))
        self.assertEqual(rebuilt["001"][0]["Chosen Price"], 9)
        self.assertEqual(rebuilt["missing"], [])


if __name__ == "__main__":
    unittest.main()
