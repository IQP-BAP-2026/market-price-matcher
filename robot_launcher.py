"""
Runs price_robot.py with the options chosen in the desktop window (price_robot_ui.py).

The window starts this script with price_robot.py's normal command-line options, plus these
environment variables:
  BAP_ROBOT_SETTINGS  JSON of setting overrides from "Advanced settings",
                      e.g. {"BAP_MARKET_PRICE_PERCENT": 0.25, "PRICE_CHANGE_REVIEW_THRESHOLD": 0.3}.
  BAP_ROBOT_FILTER    JSON {"codes": [...]}: only search these product codes
                      ("only the problems from a previous search").
  BAP_ROBOT_PERFORMANCE_FILE
                      Full path for the performance analysis workbook.
  BAP_ROBOT_CURRENT_PRICES
                      Current-price file to compare against. Empty = no comparison.
                      Not set = the robot's normal behaviour (current_prices.xlsx in the folder).
  BAP_ROBOT_PERFORMANCE
                      "0" = do not create the performance analysis workbook.

Nothing in matcher_core.py or price_robot.py is modified on disk: settings are changed only
in memory for this one run, and the code filter gives price_robot.py a temporary copy of the
products workbook that contains only the chosen rows (at their original row numbers, so
"Spreadsheet Row" in the results still points to the real file).

It also works by hand, e.g.:
  py robot_launcher.py --limit 20 --stores super99,rey
"""
from __future__ import annotations

import atexit
import json
import os
import sys
import tempfile
from copy import copy
from pathlib import Path

from app_runtime import DATA_DIR

BASE_DIR = DATA_DIR

MATCHER_KEYS = {
    "BAP_MARKET_PRICE_PERCENT", "PRICE_BASIS", "MIN_GOOD_MATCHES", "RAW_MATCH_STOP_COUNT",
    "MIN_SIZE_RATIO", "MAX_SIZE_RATIO", "PREFERRED_SIZE_RATIO_MIN", "PREFERRED_SIZE_RATIO_MAX",
    "MIN_PREFERRED_SIZE_MATCHES", "SIZE_CLUSTER_FACTOR", "MAX_FALLBACK_SIZE_MATCHES",
    "HARD_OUTLIER_LOW_RATIO", "HARD_OUTLIER_HIGH_RATIO", "LOG_MAD_Z_THRESHOLD",
    "PRICE_RATIO_WARNING", "PRICE_CV_WARNING", "MARKET_PRICE_LEVEL",
    "DUPLICATE_SIZE_TOLERANCE", "DUPLICATE_PRICE_TOLERANCE", "DUPLICATE_TOKEN_JACCARD",
}
ROBOT_KEYS = {"CACHE_TTL_HOURS"}
# Settings that belong to this launcher (the robot itself has no such option).
from review_config import PRICE_CHANGE_REVIEW_THRESHOLD
LAUNCHER_SETTINGS = {"PRICE_CHANGE_REVIEW_THRESHOLD": PRICE_CHANGE_REVIEW_THRESHOLD}


def setup_output() -> None:
    # Make sure accented product names print correctly when piped to the UI.
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)
        except Exception:
            pass


# ---------------------------------------------------------------------------
# Settings overrides
# ---------------------------------------------------------------------------
def apply_overrides() -> None:
    raw = os.environ.get("BAP_ROBOT_SETTINGS", "").strip()
    if not raw:
        return
    overrides = json.loads(raw)

    import matcher_core as mc
    import price_robot as pr

    changed = []
    for key, value in overrides.items():
        if key in LAUNCHER_SETTINGS:
            value = float(value)
            if value != LAUNCHER_SETTINGS[key]:
                changed.append(f"{key}={value} (default {LAUNCHER_SETTINGS[key]})")
            LAUNCHER_SETTINGS[key] = value
            continue
        if key in MATCHER_KEYS:
            module = mc
        elif key in ROBOT_KEYS:
            module = pr
        else:
            print(f"[SETTINGS] Ignoring unknown setting: {key}")
            continue
        default = getattr(module, key)
        # Keep the same type as the value in the code (int stays int, etc.).
        if isinstance(default, bool):
            value = bool(value)
        elif isinstance(default, int):
            value = int(value)
        elif isinstance(default, float):
            value = float(value)
        else:
            value = str(value)
        if value != default:
            changed.append(f"{key}={value} (default {default})")
        setattr(module, key, value)

    if changed:
        print("[SETTINGS] Custom settings for this run:")
        for line in changed:
            print(f"  {line}")
    else:
        print("[SETTINGS] Using default settings.")


# ---------------------------------------------------------------------------
# Reading the products workbook
# ---------------------------------------------------------------------------
def open_products(path: Path):
    """Return (workbook, sheet, header_row, header_map) using price_robot's own rules."""
    from openpyxl import load_workbook
    import price_robot as pr

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = pr.choose_sheet(wb)
    hr = pr.find_header_row(ws)
    return wb, ws, hr, pr.header_map(ws, hr)


def cell(values, col):
    if not col or col - 1 >= len(values):
        return None
    return values[col - 1]


def text(value) -> str:
    return "" if value is None else str(value).strip()


# ---------------------------------------------------------------------------
# Product-code filter ("only the problems from a previous search")
# ---------------------------------------------------------------------------
def apply_product_filter() -> None:
    raw = os.environ.get("BAP_ROBOT_FILTER", "").strip()
    if not raw:
        return
    codes = json.loads(raw).get("codes")
    if codes is None:
        return

    from openpyxl import Workbook
    import price_robot as pr

    argv = sys.argv
    if "--products" in argv:
        i = argv.index("--products") + 1
        src = Path(argv[i])
    else:
        argv += ["--products", ""]
        i = len(argv) - 1
        src = pr.PRODUCTS_FILE
    if not src.is_absolute():
        src = BASE_DIR / src

    code_set = {pr.normalize_product_code(c) for c in codes}

    wb_in, ws, hr, hm = open_products(src)
    code_col, name_col = hm.get("codigo de producto"), hm.get("nombre del producto")

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = ws.title  # keep the MASTER-YYYY name if there is one
    for r, values in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
        if r < hr:
            continue
        if r > hr:
            if not text(cell(values, name_col)):
                continue
            if pr.normalize_product_code(cell(values, code_col)) not in code_set:
                continue
        # write at the same row number as the original file
        for c, v in enumerate(values, start=1):
            if v is not None:
                ws_out.cell(row=r, column=c, value=v)
    total = sum(1 for v in ws.iter_rows(min_row=hr + 1, values_only=True) if text(cell(v, name_col)))
    wb_in.close()
    kept = sum(1 for row in ws_out.iter_rows(min_row=hr + 1, values_only=True) if text(cell(row, name_col)))

    fd, tmp = tempfile.mkstemp(prefix="bap_products_", suffix=".xlsx")
    os.close(fd)
    wb_out.save(tmp)
    atexit.register(lambda: os.path.exists(tmp) and os.remove(tmp))
    argv[i] = tmp

    print(f"[FILTER] {kept} of {total} products selected ({len(code_set)} product code(s) from last results) "
          f"from {src.name}")


def _read_current_prices(path: Path) -> dict:
    from current_prices import read_current_prices
    return read_current_prices(path)


def _output_path() -> Path:
    import price_robot as pr
    out = str(pr.OUTPUT_DIR / f"{pr.OUTPUT_BASENAME}.xlsx")
    if "--output" in sys.argv:
        out = sys.argv[sys.argv.index("--output") + 1]
    path = Path(out)
    return path if path.is_absolute() else pr.BASE_DIR / path


def _mark_performance_workbook(wb, comparison: dict, threshold: float) -> None:
    """Mark large price changes for manual review in the performance analysis (open workbook)."""
    from openpyxl.styles import Font, PatternFill
    import results_format as rf
    ws = rf.find_sheet(wb, rf.ANALYSIS_SHEET)
    if ws is None:
        return
    headers = [rf.internal_header(c.value) for c in ws[1]]
    if "Needs Manual Review" not in headers:
        return
    i_review = headers.index("Needs Manual Review") + 1
    i_row = headers.index("Spreadsheet Row") + 1
    i_flag = headers.index("Quality Flag") + 1
    reason_col = ws.max_column + 1
    head = ws.cell(1, reason_col, rf.header_es("Review Reason"))
    head.fill, head.font = copy(ws.cell(1, 1).fill), copy(ws.cell(1, 1).font)
    head.alignment = copy(ws.cell(1, 1).alignment)
    ws.column_dimensions[head.column_letter].width = 42
    flag_fill = PatternFill("solid", fgColor="FDE2C4")
    for r in range(2, ws.max_row + 1):
        reasons = []
        quality = str(ws.cell(r, i_flag).value or "")
        if quality != "OK":
            reasons.append(rf.alerts_es(quality))
        _cur, diff, big, _id = comparison.get(str(ws.cell(r, i_row).value).strip(), (None, None, "", ""))
        if big == "YES":
            reasons.append(f"Cambio de precio de {diff:+.0%} frente al actual (límite ±{threshold:.0%})")
            cell = ws.cell(r, i_review, rf.value_es("Needs Manual Review", "YES"))
            cell.fill = flag_fill
            cell.font = Font(bold=True, color="9A3412")
        ws.cell(r, reason_col, " · ".join(x for x in reasons if x))


def check_current_prices_file() -> None:
    """Stop before searching when the chosen current-prices file lacks required columns or can't be read."""
    chosen = os.environ.get("BAP_ROBOT_CURRENT_PRICES", "").strip()
    if not chosen:
        return
    from current_prices import current_prices_columns_missing, read_current_rows
    missing = current_prices_columns_missing(chosen)
    if missing:
        raise SystemExit(f"[ERROR] The current-prices file is missing required column(s): {', '.join(missing)}")
    rows = read_current_rows(chosen)   # raises a readable error for duplicate codes, bad prices, …
    if rows and not any(row.get("Id") for row in rows):
        raise SystemExit("[ERROR] The current-prices file has an Id column but no Salesforce Ids in it.")


def apply_current_prices_choice() -> None:
    """Point the robot's price comparison at the file chosen in the UI (or switch it off),
    and flag products whose estimate differs a lot from the current BAP price."""
    if "BAP_ROBOT_CURRENT_PRICES" not in os.environ:
        return
    import price_robot as pr

    chosen = os.environ["BAP_ROBOT_CURRENT_PRICES"].strip()
    # A path that never exists makes both robot functions skip the comparison cleanly.
    target = chosen or str(BASE_DIR / "__price_comparison_off__.xlsx")
    original_compare = pr.compare_current_prices
    original_analysis = pr.create_performance_analysis
    threshold = LAUNCHER_SETTINGS["PRICE_CHANGE_REVIEW_THRESHOLD"]
    comparison: dict = {}   # "Spreadsheet Row" -> (current price, % difference, "YES"/"", Salesforce Id)
    state = {"added": False}

    def build_comparison(summary_headers, summary_rows) -> None:
        if comparison:
            return
        from current_prices import read_current_rows
        current_rows = read_current_rows(Path(target))
        prices = {r["Código de producto"]: r["Precio de lista"] for r in current_rows if r["Precio de lista"] is not None}
        ids = {r["Código de producto"]: r.get("Id") or "" for r in current_rows}
        hi = {h: i for i, h in enumerate(summary_headers)}
        for row in summary_rows:
            key = pr.normalize_product_code(row[hi["Código de producto"]])
            cur = prices.get(key)
            est = pr.parse_money(row[hi["Estimated New Product Price"]])
            diff = None if cur in (None, 0) or est is None else (est - cur) / cur
            flag = "YES" if diff is not None and abs(diff) > threshold else ""
            comparison[str(row[hi["Spreadsheet Row"]]).strip()] = (cur, diff, flag, ids.get(key, ""))

    def summary_hook(summary_headers, summary_rows):
        # Add each product's current price, the % difference, the big-change mark and its Salesforce Id
        # (so the results file can be reviewed and exported on its own) before anything is written.
        try:
            build_comparison(summary_headers, summary_rows)
            import results_format as rf
            rf.CHANGE_LIMIT = threshold
            i_row = summary_headers.index("Spreadsheet Row")
            summary_headers.extend(["Current BAP Price", "% Difference vs Current", "Large Price Change", "Salesforce Id"])
            for row in summary_rows:
                cur, diff, big, sf_id = comparison.get(str(row[i_row]).strip(), (None, None, "", ""))
                row.extend([cur, diff, big, sf_id or None])
            state["added"] = True
        except Exception as exc:
            print(f"[PRICE CHANGE] Could not add the price comparison to the results: {exc!r}")

    def compare(summary_rows, summary_headers, current_prices_file=None):
        if not chosen:
            print("\n[PRICE COMPARISON] Off (not selected in the window).")
            return None
        result = original_compare(summary_rows, summary_headers, current_prices_file=target)
        try:
            build_comparison(summary_headers, summary_rows)
            big = sum(1 for v in comparison.values() if v[2] == "YES")
            print(f"[PRICE CHANGE] {big} product(s) differ from the current BAP price by more than "
                  f"±{threshold:.0%} and are marked for review.")
            if state["added"]:
                print(f"[PRICE CHANGE] Added current price and % difference to {_output_path().name}")
        except Exception as exc:
            print(f"[PRICE CHANGE] Could not add the price comparison to the results: {exc!r}")
        return result

    def performance_hook(wb):
        if chosen and comparison:
            try:
                _mark_performance_workbook(wb, comparison, threshold)
            except Exception as exc:
                print(f"[PRICE CHANGE] Could not mark price changes in the analysis: {exc!r}")

    def analysis(*args, current_prices_file=None, **kwargs):
        return original_analysis(*args, current_prices_file=target, **kwargs)

    if chosen:
        pr.SUMMARY_HOOKS.append(summary_hook)
        pr.PERFORMANCE_WORKBOOK_HOOKS.append(performance_hook)
    pr.compare_current_prices = compare
    pr.create_performance_analysis = analysis
    print(f"[PRICE COMPARISON] {'Using ' + chosen if chosen else 'Off'}")


def apply_performance_choice() -> None:
    """Skip the performance analysis workbook when it is switched off in the window."""
    if os.environ.get("BAP_ROBOT_PERFORMANCE", "").strip() != "0":
        return
    import price_robot as pr

    def skipped(*args, **kwargs):
        print("[PERFORMANCE ANALYSIS] Skipped (not selected in the window).")

    pr.create_performance_analysis = skipped


def main() -> None:
    setup_output()
    os.chdir(BASE_DIR)  # current_prices.xlsx is looked up relative to the working folder

    apply_overrides()
    perf_file = os.environ.get("BAP_ROBOT_PERFORMANCE_FILE", "").strip()
    if perf_file:  # name/folder chosen in the window; BASE_DIR / an absolute path = that path
        import price_robot as pr
        pr.PERFORMANCE_BASENAME = perf_file
    check_current_prices_file()
    apply_current_prices_choice()
    apply_performance_choice()
    apply_product_filter()
    import price_robot
    price_robot.main()


if __name__ == "__main__":
    main()
