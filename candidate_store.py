"""Disposable, fingerprinted candidate index: load one product's evidence at a time."""
from collections import Counter
from contextlib import closing
from functools import lru_cache
import hashlib
from itertools import groupby
import json
from pathlib import Path
import sqlite3
import tempfile
import zlib
from app_runtime import cache_file


def file_hash(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def product_code(value):
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value or "").strip()


def build_candidate_index(results, headers, rows, fingerprint=None):
    results = Path(results)
    destination = cache_file(results.name + ".candidates.sqlite3", results.parent)
    code_index = headers.index("Código de producto")
    with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".candidates.tmp", delete=False) as temp:
        staging = Path(temp.name)
    try:
        with closing(sqlite3.connect(staging)) as db:
            db.execute("CREATE TABLE metadata (fingerprint TEXT, version INTEGER)")
            db.execute("CREATE TABLE products (code TEXT PRIMARY KEY, candidates BLOB, counts TEXT)")
            for key, group in groupby(rows, key=lambda row: product_code(row[code_index])):
                candidates = [dict(zip(headers, row)) for row in group]
                previous = db.execute("SELECT candidates FROM products WHERE code=?", (key,)).fetchone()
                if previous:
                    candidates = json.loads(zlib.decompress(previous[0])) + candidates
                counts = Counter(str(c.get("Store") or c.get("Store Name")) for c in candidates if c.get("Status") == "ACCEPTED")
                db.execute("INSERT OR REPLACE INTO products VALUES (?, ?, ?)",
                           (key, zlib.compress(json.dumps(candidates, ensure_ascii=False, separators=(",", ":"), default=str).encode("utf-8"), 1), json.dumps(counts)))
            db.execute("INSERT INTO metadata VALUES (?, 1)", (fingerprint or file_hash(results),))
            db.commit()
        staging.replace(destination)
    finally:
        staging.unlink(missing_ok=True)
    return destination


class CandidateStore:
    def __init__(self, results, fingerprint):
        results = Path(results)
        self.path = cache_file(results.name + ".candidates.sqlite3", results.parent)
        self.fingerprint = fingerprint
        if not self._valid():
            from openpyxl import load_workbook
            wb = load_workbook(results, read_only=True, data_only=True)
            try:
                if "Candidates" not in wb.sheetnames:
                    self.counts = {}
                    self.path = None
                    return
                rows = wb["Candidates"].iter_rows(values_only=True)
                headers = list(next(rows))
                build_candidate_index(results, headers, rows, fingerprint)
            finally:
                wb.close()
        with closing(sqlite3.connect(self.path)) as db:
            self.counts = {key: json.loads(counts) for key, counts in db.execute("SELECT code, counts FROM products")}

    def _valid(self):
        if not self.path.is_file():
            return False
        try:
            with closing(sqlite3.connect(self.path)) as db:
                return db.execute("SELECT fingerprint, version FROM metadata").fetchone() == (self.fingerprint, 1)
        except sqlite3.DatabaseError:
            return False

    @lru_cache(maxsize=32)
    def __getitem__(self, key):
        if self.path is None:
            return []
        with closing(sqlite3.connect(self.path)) as db:
            if db.execute("SELECT fingerprint FROM metadata").fetchone() != (self.fingerprint,):
                raise ValueError("Results changed during review. Reopen the review before choosing candidates.")
            row = db.execute("SELECT candidates FROM products WHERE code=?", (key,)).fetchone()
        return json.loads(zlib.decompress(row[0])) if row else []
