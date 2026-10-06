"""Shared header-based current-price input for comparison and human review."""
import csv
import io
import math
from pathlib import Path

# Excel and Salesforce save CSV files in different ways: UTF-8 (with or without a BOM), "Unicode text"
# (UTF-16), or the Windows code page (cp1252) — where "ñ" and accented letters are single bytes that
# aren't valid UTF-8 ("invalid continuation byte"). Spanish-language Excel also separates columns with
# ";" and writes decimals with a comma (5,02).
_TEXT_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")


def is_spreadsheet(path):
    """True for an Excel workbook (.xlsx/.xlsm are ZIP files), whatever the file is called. An old .xls
    workbook can't be read: ask for it to be saved as .xlsx or .csv."""
    with open(path, "rb") as stream:
        start = stream.read(8)
    if start.startswith(b"\xd0\xcf\x11\xe0"):
        raise ValueError(f"{Path(path).name} is an old Excel (.xls) file. Open it in Excel and save it as "
                         ".xlsx (Excel workbook) or .csv, then choose it again.")
    return start.startswith(b"PK")


def read_text_table(path):
    """(rows, decimal_comma) of a CSV/text file in any of the encodings and separators Excel uses."""
    raw = Path(path).read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = raw.decode("utf-16")
    else:
        for encoding in _TEXT_ENCODINGS:
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text, newline=""), dialect))
    return rows, dialect.delimiter == ";"


def product_code(value):
    if value is None:
        return ""
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()


def money(value, decimal_comma=False):
    """A price from a cell: 5.02, "$5.02", "1,234.50" — or, with decimal_comma (Spanish-format CSV),
    "5,02" and "1.234,50"."""
    try:
        if isinstance(value, bool):
            return None
        if isinstance(value, (int, float)):
            result = float(value)
        else:
            text = str(value).replace("$", "").replace("\u00a0", "").replace(" ", "").strip()
            if "," in text and "." in text:      # both: the last one is the decimal separator
                comma_decimal = text.rfind(",") > text.rfind(".")
            else:                                  # only a comma: decimal in Spanish-format files
                comma_decimal = decimal_comma and "," in text
            text = text.replace(".", "").replace(",", ".") if comma_decimal else text.replace(",", "")
            result = float(text)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def read_current_rows(path):
    """Read Salesforce CSV/XLSX or legacy BAP headers; never infer price by position.

    Preserve source columns and add canonical BAP names for existing consumers.
    IDs are PricebookEntry IDs from the input, never product codes or product IDs.
    """
    path = Path(path)
    decimal_comma = False
    if not is_spreadsheet(path):   # CSV / text, whatever its extension
        rows, decimal_comma = read_text_table(path)
        headers, values = (rows[0], rows[1:]) if rows else ([], [])
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            iterator = wb.active.iter_rows(values_only=True)
            headers = next(iterator, ())
            values = list(iterator)
        finally:
            wb.close()
    headers = [str(h or "").strip() for h in headers]
    if len(set(h.casefold() for h in headers if h)) != len([h for h in headers if h]):
        raise ValueError("Duplicate column headers in current-prices file.")
    fields = {h.casefold(): h for h in headers}
    code_field = fields.get("productcode") or fields.get("código de producto")
    price_field = fields.get("unitprice") or fields.get("precio de lista")
    if not code_field or not price_field:
        raise ValueError("Current prices require ProductCode and UnitPrice headers (legacy Código de producto / Precio de lista also supported).")
    name_field = fields.get("name") or fields.get("nombre del producto")
    id_field = fields.get("id")
    rows, codes, ids = [], set(), set()
    for row_number, values_row in enumerate(values, start=2):
        if not any(value not in (None, "") for value in values_row):
            continue
        row = dict(zip(headers, values_row))
        key = product_code(row.get(code_field))
        if not key:
            raise ValueError(f"Current-prices row {row_number} is missing ProductCode.")
        if key in codes:
            raise ValueError(f"Multiple current-price entries for ProductCode {key}. Filter the input to one pricebook entry per product.")
        codes.add(key)
        identifier = str(row.get(id_field) or "").strip()
        if identifier and identifier in ids:
            raise ValueError(f"Duplicate Salesforce Id in current prices: {identifier}")
        if identifier:
            ids.add(identifier)
        raw_price = row.get(price_field)
        price = money(raw_price, decimal_comma)
        if raw_price not in (None, "") and (price is None or price < 0):
            raise ValueError(f"Invalid UnitPrice for ProductCode {key}: {raw_price}")
        row.update({"Código de producto": key, "Nombre del producto": row.get(name_field),
                    "Precio de lista": price, "Id": identifier})
        rows.append(row)
    return rows


def read_current_prices(path):
    return {row["Código de producto"]: row["Precio de lista"] for row in read_current_rows(path)
            if row["Precio de lista"] is not None}


# ---------------------------------------------------------------------------
# Required columns — checked before a search starts (window and command line)
# ---------------------------------------------------------------------------
# Every column of the products sheet that the matcher reads.
PRODUCT_REQUIRED_COLUMNS = (
    "Código de producto", "Nombre del producto", "Unidad de Peso en KG", "Sub-familia de Productos",
    "Linea de producto", "Familia de productos", "Descripción del producto", "Grupo de Productos",
    "Categoria RepTrim",
)
# Current prices must be the Salesforce pricebook download: the reviewed
# prices are exported by Salesforce Id, so a file without the Id column (like the old BAP current_prices.xlsx)
# is refused. (Salesforce name, older BAP name) — either name is accepted for code and price.
CURRENT_REQUIRED_COLUMNS = (("ProductCode", "Código de producto"), ("Id",), ("UnitPrice", "Precio de lista"))


def _norm(value):
    from matcher_core import normalize
    return normalize(value)


def product_columns_missing(path):
    """Required columns missing from the products workbook ([] = fine). The robot reads the sheet/header row
    that has the most of them (a MASTER-yyyy sheet first), like it does when searching."""
    import re
    from openpyxl import load_workbook
    wanted = {_norm(c): c for c in PRODUCT_REQUIRED_COLUMNS}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = [s for s in wb.sheetnames if re.fullmatch(r"MASTER-\d{4}", s, re.I)]
        names += [s for s in wb.sheetnames if s not in names]
        best = set()
        for name in names:
            for row in wb[name].iter_rows(max_row=30, values_only=True):
                found = {_norm(v) for v in row if v not in (None, "")} & set(wanted)
                if len(found) > len(best):
                    best = found
            if len(best) == len(wanted):
                break
    finally:
        wb.close()
    return [wanted[k] for k in wanted if k not in best]


def current_prices_columns_missing(path):
    """Required columns missing from the current-prices file ([] = fine), as 'ProductCode / Código de producto'."""
    path = Path(path)
    if not is_spreadsheet(path):
        rows, _ = read_text_table(path)
        headers = rows[0] if rows else []
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            headers = next(wb.active.iter_rows(max_row=1, values_only=True), ())
        finally:
            wb.close()
    present = {str(h or "").strip().casefold() for h in headers}
    return [" / ".join(pair) for pair in CURRENT_REQUIRED_COLUMNS
            if not any(name.casefold() in present for name in pair)]
