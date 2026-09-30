from __future__ import annotations

import argparse
import math
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import pdfplumber
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# One branch from each requested supermarket, all on page 1 of the supplied ACODECO PDF.
# X coordinates are the centers of the four table columns on that page.
SELECTED_STORES = {
    "Rey - Via España": 244.8,
    "Riba Smith - Bella Vista": 289.2,
    "Super 99 - Camino Real": 479.1,
    "Xtra - Transistmica": 390.0,
}
PRICE_X_TOLERANCE = 18.0


def norm(s: object) -> str:
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    s = s.lower().replace("×", "x")
    s = re.sub(r"\bkilogramos?\b", "kg", s)
    s = re.sub(r"\bgramos?\b", "g", s)
    s = re.sub(r"\bmililitros?\b", "ml", s)
    s = re.sub(r"\blitros?\b", "l", s)
    s = re.sub(r"\bonzas?\b", "oz", s)
    s = re.sub(r"\blibras?\b", "lb", s)
    s = re.sub(r"\bunidades?\b", "unidad", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


@dataclass(frozen=True)
class Qty:
    dimension: str  # mass, volume, count
    amount: float   # g, ml, or units
    source: str


MASS = {"g": 1.0, "kg": 1000.0, "oz": 28.349523125, "lb": 453.59237}
VOLUME = {"ml": 1.0, "l": 1000.0, "lt": 1000.0, "gal": 3785.411784, "gl": 3785.411784}


def _qty(value: float, unit: str, source: str) -> Optional[Qty]:
    unit = unit.lower()
    if unit in MASS:
        return Qty("mass", value * MASS[unit], source)
    if unit in VOLUME:
        return Qty("volume", value * VOLUME[unit], source)
    if unit in {"unidad", "unid", "und"}:
        return Qty("count", value, source)
    if unit in {"docena", "docenas"}:
        return Qty("count", value * 12.0, source)
    return None


def parse_measure(text: str) -> Optional[Qty]:
    """Parse ACODECO measure, e.g. '1,42 Litros', '155 Gramos', '1 Docena'."""
    t = norm(text).replace(",", ".")
    m = re.search(r"(\d+(?:\.\d+)?)\s*(kg|g|ml|l|oz|lb|docena|unidad)\b", t)
    return _qty(float(m.group(1)), m.group(2), "pdf_measure") if m else None


def parse_package_quantity(name: str, code: str = "") -> Optional[Qty]:
    """Parse total package size. Multipacks are multiplied: 12x170g -> 2040g."""
    # Prefer the human-readable product name; use code only as fallback.
    for source_name, raw in (("product_name", name), ("product_code", code)):
        t = norm(raw).replace(",", ".")

        # 12 x 170 g, 6x2.5lt, 30x2lb, 24 x 400 ml, etc.
        m = re.search(r"\b(\d+)\s*x\s*(\d+(?:\.\d+)?)\s*(kg|g|ml|l|lt|oz|lb|gal|gl)\b", t)
        if m:
            q = _qty(float(m.group(2)), m.group(3), source_name + ":multipack")
            return Qty(q.dimension, int(m.group(1)) * q.amount, q.source) if q else None

        # 5 docenas / 3 docenas
        m = re.search(r"\b(\d+(?:\.\d+)?)\s*docenas?\b", t)
        if m:
            return Qty("count", float(m.group(1)) * 12.0, source_name + ":dozens")

        # 30 unidades / 10 unid
        m = re.search(r"\b(\d+(?:\.\d+)?)\s*(?:unidades?|unid|und)\b", t)
        if m:
            return Qty("count", float(m.group(1)), source_name + ":count")

        # Single mass/volume. Last occurrence is often the actual pack size after descriptors.
        matches = list(re.finditer(r"\b(\d+(?:\.\d+)?)\s*(kg|g|ml|l|lt|oz|lb|gal|gl)\b", t))
        if matches:
            m = matches[-1]
            return _qty(float(m.group(1)), m.group(2), source_name + ":single")
    return None


# Conservative identity rules. A row is matched only when one and only one rule accepts it.
# Generic/varied products are deliberately rejected when the PDF has a narrower subtype.
# Add rules here as the spreadsheet evolves; no LLM/fuzzy matching is used.
RULES = {
    "Sardina en Salsa de Tomate sin Picante": dict(any=[r"sardina.*salsa de tomate.*sin picante"], forbid=[r"picante(?!.*sin picante)", r"aceite", r"mostaza"]),
    "Tuna en Agua": dict(any=[r"\batun\b.*\bagua\b", r"\btuna\b.*\bagua\b"], forbid=[r"aceite"]),
    "Arroz de Primera": dict(any=[r"arroz de primera"], forbid=[r"integral", r"molido", r"precoc", r"jazmin", r"risotto"]),
    "Codito": dict(any=[r"\bcodito(?:s)?\b"], forbid=[r"variad", r"queso"]),
    "Hojuelas de Maíz (corn flakes) cajeta": dict(any=[r"hojuelas? de maiz", r"corn flakes"], forbid=[]),
    "Macarrones (espaguetti) Cal 3- 5": dict(any=[r"\bmacarrones?\b", r"\bespaguet(?:ti|is)?\b"], forbid=[r"queso", r"variad"]),
    "Pan Molde": dict(any=[r"pan (?:de )?molde"], forbid=[]),
    "Tortilla de Maíz": dict(any=[r"tortilla(?:s)? de maiz"], forbid=[r"masa"]),
    "Ajo": dict(any=[r"\bajo\b"], forbid=[r"puerro", r"mantequilla", r"salsa", r"hummus"]),
    "Cebolla Amarilla": dict(any=[r"cebolla amarilla"], forbid=[r"caramel", r"aros"]),
    "Lechuga": dict(any=[r"\blechuga\b"], forbid=[r"romana"]),
    "Ñame Diamante": dict(any=[r"name diamante"], forbid=[]),
    "Papa Nacional": dict(any=[r"papa nacional"], forbid=[r"congel", r"prefrit", r"pre-frit", r"rayad", r"pure"]),
    "Plátano Verde": dict(any=[r"platano verde"], forbid=[r"crema", r"congel", r"maduro"]),
    "Repollo Verde": dict(any=[r"repollo verde"], forbid=[r"chino"]),
    "Tomate Nacional (Perita ó 3 x 3)": dict(any=[r"tomate (?:nacional|perita)"], forbid=[r"cherry", r"pasta", r"salsa", r"base", r"picado"]),
    "Yuca": dict(any=[r"\byuca\b"], forbid=[r"congel", r"pan"]),
    "Zanahoria": dict(any=[r"\bzanahoria(?:s)?\b"], forbid=[]),
    "Frijoles Chiricanos": dict(any=[r"frijoles? chiricanos?"], forbid=[r"molido", r"lata"]),
    "Lentejas": dict(any=[r"\blentejas?\b"], forbid=[r"menestras? variad"]),
    "Porotos Rojos": dict(any=[r"porotos? rojos?", r"frijoles? rojos?"], forbid=[r"molido", r"lata", r"menestras? variad"]),
    "Guineos": dict(any=[r"\bguineos?\b"], forbid=[]),
    "Manzana Roja Mediana": dict(any=[r"manzana roja mediana"], forbid=[]),
    "Naranja de Jugo": dict(any=[r"naranja de jugo"], forbid=[r"jugo", r"concentr"]),
    "Piña": dict(any=[r"\bpina\b"], forbid=[r"lata", r"jugo", r"conserv"]),
    "Margarina en barra": dict(any=[r"margarina.*barra"], forbid=[]),
    "Leche Evaporada": dict(any=[r"leche evaporada"], forbid=[]),
    "Leche Fresca y Pasteurizada": dict(any=[r"leche (?:fresca|pasteurizada)"], forbid=[r"polvo", r"evaporada", r"soya", r"almendra", r"sabor"]),
    "Leche en Polvo entera instantánea": dict(any=[r"leche en polvo.*(?:entera|instantanea)", r"leche.*entera.*polvo"], forbid=[r"infant", r"bebe", r"diabet", r"adulto", r"desayuno"]),
    "Queso Amarillo": dict(any=[r"queso amarillo"], forbid=[]),
    "Queso Blanco Prensado Bajo en Sal": dict(any=[r"queso blanco.*prensado.*bajo.*sal"], forbid=[]),
    "Huevos Medianos de Gallina": dict(any=[r"huevos? medianos? de gallina"], forbid=[r"congel"]),
    "Azúcar Morena": dict(any=[r"azucar morena"], forbid=[r"dieta", r"micropulver"]),
    "Café Molido Tradicional": dict(any=[r"cafe molido tradicional"], forbid=[r"filtro"]),
    "Jugo de Naranja": dict(any=[r"jugo de naranja"], forbid=[r"concentr"]),
    "Mayonesa": dict(any=[r"\bmayonesa\b"], forbid=[]),
    "Pasta de Tomate": dict(any=[r"pasta de tomate"], forbid=[]),
    "Sal": dict(any=[r"^sal(?:\s|$)"], forbid=[r"salsa"]),
    "Salsa de Tomate": dict(any=[r"salsa de tomate"], forbid=[r"casera"]),
    "Soda en envase plástico": dict(any=[r"\bsoda\b.*(?:plastico|plastica|envase plastico)"], forbid=[r"lata", r"vidrio"]),
    "Sopa Deshidratada (Pollo y Fideos)": dict(any=[r"sopa.*pollo.*fideos"], forbid=[r"liquida", r"lata"]),
    "Te Negro": dict(any=[r"te negro"], forbid=[r"frio"]),
}


def identity_candidates(product_name: str, description: str = "") -> list[str]:
    t = norm(product_name + " " + (description or ""))
    out = []
    for pdf_name, rule in RULES.items():
        if any(re.search(p, t) for p in rule["forbid"]):
            continue
        if any(re.search(p, t) for p in rule["any"]):
            out.append(pdf_name)
    return out


def extract_pdf_catalog(pdf_path: Path):
    """Extract page-1 product, measure and selected store prices by table geometry."""
    with pdfplumber.open(pdf_path) as pdf:
        page = pdf.pages[0]
        words = page.extract_words(x_tolerance=1, y_tolerance=1)

    # Product labels define row baselines. Prices are typically ~0.3 pt above label text,
    # so collect all words within 0.8 pt of each label baseline rather than exact grouping.
    anchors = sorted({round(w["top"], 1) for w in words if w["x0"] < 170 and 76 <= w["top"] <= 575})
    catalog = {}
    for top in anchors:
        ws = [w for w in words if abs(w["top"] - top) <= 0.8]
        left = sorted([w for w in ws if w["x0"] < 170], key=lambda w: w["x0"])
        measure_words = sorted([w for w in ws if 170 <= w["x0"] < 225], key=lambda w: w["x0"])
        if not left or not measure_words:
            continue
        product = " ".join(w["text"] for w in left).strip()
        measure = " ".join(w["text"] for w in measure_words).strip()
        if product in {"Producto", "Nota.", "Fuente:"}:
            continue
        prices = {}
        for store, center in SELECTED_STORES.items():
            hits = [w for w in ws if abs(((w["x0"] + w["x1"]) / 2) - center) <= PRICE_X_TOLERANCE]
            nums = []
            for w in hits:
                try:
                    nums.append(float(w["text"].replace(",", ".")))
                except ValueError:
                    pass
            prices[store] = nums[0] if len(nums) == 1 else None
        catalog[product] = {"measure": measure, "qty": parse_measure(measure), "prices": prices}
    return catalog


def find_pdf_key(catalog: dict, canonical_name: str) -> Optional[str]:
    n = norm(canonical_name)
    exact = [k for k in catalog if norm(k) == n]
    return exact[0] if len(exact) == 1 else None


def scaled_price(base_price: Optional[float], pdf_qty: Qty, target_qty: Qty) -> Optional[float]:
    if base_price is None or pdf_qty.dimension != target_qty.dimension or pdf_qty.amount <= 0:
        return None
    return base_price * target_qty.amount / pdf_qty.amount


def main():
    ap = argparse.ArgumentParser(description="Strict deterministic matcher: Banco de Alimentos product spreadsheet -> ACODECO supermarket PDF")
    ap.add_argument(
        "xlsx", nargs="?", type=Path, default=Path("products.xlsx"),
        help="Product spreadsheet (default: products.xlsx in the current directory)",
    )
    ap.add_argument(
        "pdf", nargs="?", type=Path, default=Path("acodeco_dataset.pdf"),
        help="ACODECO price PDF (default: acodeco_dataset.pdf in the current directory)",
    )
    ap.add_argument("-o", "--output", type=Path, default=Path("matched_prices.xlsx"))
    args = ap.parse_args()

    if not args.xlsx.is_file():
        raise FileNotFoundError(
            f"Could not find product spreadsheet: {args.xlsx.resolve()}\n"
            "Place products.xlsx in the same directory where you run this script, "
            "or pass a different spreadsheet path as the first argument."
        )
    if not args.pdf.is_file():
        raise FileNotFoundError(
            f"Could not find price PDF: {args.pdf.resolve()}\n"
            "Place acodeco_dataset.pdf in the same directory where you run this script, "
            "or pass a different PDF path as the second argument."
        )

    catalog = extract_pdf_catalog(args.pdf)
    if len(catalog) < 50:
        raise RuntimeError(f"PDF parse looks incomplete: only {len(catalog)} product rows found")

    src = load_workbook(args.xlsx, data_only=True, read_only=True)
    ws = src[src.sheetnames[0]]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}
    required = ["Nombre del producto", "Código de producto", "Unidad de Peso en KG", "Unidad de Medida", "Descripción del producto"]
    missing = [h for h in required if h not in idx]
    if missing:
        raise KeyError(f"Spreadsheet is missing columns: {missing}")

    out = Workbook()
    ows = out.active
    ows.title = "Price Matches"
    out_headers = [
        "ERP Id (QBO)", "Nombre del producto", "Código de producto", "Status", "PDF product", "PDF measure",
        "Target quantity", "Quantity basis", "Scale factor",
        *SELECTED_STORES.keys(), "Reason"
    ]
    ows.append(out_headers)

    erp_i = idx.get("ERP Id (QBO)")
    matched = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[idx["Nombre del producto"]] or "").strip()
        if not name:
            continue
        code = str(row[idx["Código de producto"]] or "").strip()
        desc = str(row[idx["Descripción del producto"]] or "").strip()
        erp = row[erp_i] if erp_i is not None else None

        candidates = identity_candidates(name, desc)
        status, reason = "UNMATCHED", "No conservative identity rule matched"
        pdf_key = None
        target_qty = parse_package_quantity(name, code)

        if len(candidates) > 1:
            reason = "Ambiguous identity: " + "; ".join(candidates)
        elif len(candidates) == 1:
            pdf_key = find_pdf_key(catalog, candidates[0])
            if pdf_key is None:
                reason = f"Expected PDF row not found exactly: {candidates[0]}"
            elif target_qty is None:
                # Conservative fallback to explicit spreadsheet kg-weight only for mass-priced PDF rows.
                pdf_qty = catalog[pdf_key]["qty"]
                weight = row[idx["Unidad de Peso en KG"]]
                try:
                    weight = float(weight)
                except (TypeError, ValueError):
                    weight = 0
                if pdf_qty and pdf_qty.dimension == "mass" and weight > 0:
                    target_qty = Qty("mass", weight * 1000.0, "Unidad de Peso en KG")
                else:
                    reason = "Product identity matched, but target package quantity could not be parsed safely"

            if pdf_key and target_qty:
                pdf_qty = catalog[pdf_key]["qty"]
                if pdf_qty is None:
                    reason = "PDF measure could not be parsed"
                elif pdf_qty.dimension != target_qty.dimension:
                    reason = f"Unit dimension mismatch: PDF={pdf_qty.dimension}, spreadsheet={target_qty.dimension}; no density/guess conversion"
                else:
                    status = "MATCHED"
                    reason = "Exact rule-based identity + compatible quantity dimension"
                    matched += 1

        if status == "MATCHED":
            info = catalog[pdf_key]
            factor = target_qty.amount / info["qty"].amount
            store_prices = [scaled_price(info["prices"][s], info["qty"], target_qty) for s in SELECTED_STORES]
            unit_label = {'mass': 'g', 'volume': 'ml', 'count': 'units'}[target_qty.dimension]
            target_text = f"{target_qty.amount:g} {unit_label}"
            ows.append([erp, name, code, status, pdf_key, info["measure"], target_text, target_qty.source, factor, *store_prices, reason])
        else:
            ows.append([erp, name, code, status, pdf_key, catalog[pdf_key]["measure"] if pdf_key else None,
                        None if target_qty is None else f"{target_qty.amount:g} {target_qty.dimension}",
                        None if target_qty is None else target_qty.source, None,
                        *([None] * len(SELECTED_STORES)), reason])

    # Formatting
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ows[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    ows.freeze_panes = "A2"
    ows.auto_filter.ref = ows.dimensions
    widths = [14, 48, 24, 13, 42, 18, 18, 24, 13, 22, 24, 24, 22, 58]
    for i, width in enumerate(widths, 1):
        ows.column_dimensions[get_column_letter(i)].width = width
    for col in range(10, 14):
        for cell in ows.iter_cols(min_col=col, max_col=col, min_row=2):
            cell[0].number_format = '$0.00'
    for cell in ows.iter_cols(min_col=9, max_col=9, min_row=2):
        cell[0].number_format = '0.0000'

    args.output.parent.mkdir(parents=True, exist_ok=True)
    out.save(args.output)
    print(f"PDF products parsed: {len(catalog)}")
    print(f"Spreadsheet rows matched: {matched}")
    print(f"Output: {args.output}")


if __name__ == "__main__":
    main()
