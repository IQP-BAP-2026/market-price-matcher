"""Explainable confidence, durable human decisions, and template-shaped exports."""
from __future__ import annotations

import csv
import json
import math
import re
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

import review_config as cfg
from current_prices import read_current_rows
from matcher_core import PRICE_CV_WARNING

CODE = "Código de producto"
NAME = "Nombre del producto"
PRICE = "Estimated New Product Price"


def number(value):
    try:
        if isinstance(value, bool):
            return None
        result = float(str(value).replace("$", "").replace(",", "").strip())
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def code(value):
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value or "").strip()


def valid_price(value):
    value = number(value)
    return value is not None and round(value, cfg.REVIEW_PRICE_DECIMALS) > 0


def table(path, sheet=None):
    if Path(path).suffix.lower() == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet in wb.sheetnames else wb.active
        rows = ws.iter_rows(values_only=True)
        headers = next(rows)
        return [dict(zip(headers, row)) for row in rows if any(v is not None for v in row)]
    finally:
        wb.close()


def confidence(row, current, counts, settings=None, lang="en"):
    settings = settings or {}
    tr = lambda en, es: es if lang == "es" else en
    threshold = float(settings.get("PRICE_CHANGE_REVIEW_THRESHOLD", cfg.PRICE_CHANGE_REVIEW_THRESHOLD))
    cv_limit = float(settings.get("PRICE_CV_WARNING", PRICE_CV_WARNING))
    estimate = number(row.get(PRICE))
    stores = int(number(row.get("Store Count")) or len(counts))
    accepted = int(number(row.get("Accepted Count")) or sum(counts.values()))
    query = str(row.get("Search Confidence") or "UNKNOWN").upper()
    notes = [tr(f"Evidence: {accepted} comparable products from {stores} stores.",
                f"Evidencia: {accepted} productos comparables de {stores} tiendas."),
             tr(f"Search-query confidence: {query}.", f"Confianza de búsqueda: {query}.")]
    if counts:
        notes.append(tr("Comparable products per store: ", "Productos comparables por tienda: ") +
                     "; ".join(f"{s}: {n}" for s, n in sorted(counts.items())))
    has_estimate = valid_price(estimate)
    if not has_estimate:
        notes.append(tr("No usable price was found. Enter a price manually or reject this item.",
                        "No se encontró un precio válido. Ingrese un precio o rechace este artículo."))
    weights = cfg.REVIEW_WEIGHTS
    points = weights["stores"] * min(1, stores / cfg.REVIEW_TARGET_STORES)
    # Capping each store prevents one large catalog from dominating the evidence.
    support = sum(min(n, cfg.REVIEW_MATCHES_PER_STORE) for n in counts.values())
    points += weights["matches"] * min(1, support / (cfg.REVIEW_TARGET_STORES * cfg.REVIEW_MATCHES_PER_STORE))
    points += weights["query"] * cfg.REVIEW_QUERY_CREDIT.get(query, 0)
    cap = 5
    if valid_price(current) and has_estimate:
        change = abs(estimate - current) / current
        within = change <= threshold + 1e-12
        points += weights["current"] * (1 if within else max(0, threshold / change))
        notes.append(tr(f"Price change: {(estimate-current)/current:+.1%}; warning limit: ±{threshold:.0%}.",
                        f"Cambio de precio: {(estimate-current)/current:+.1%}; límite: ±{threshold:.0%}."))
        if not within:
            cap = cfg.REVIEW_LARGE_CHANGE_CAP
            notes.append(tr("The change exceeds the warning limit. Verify the product and package size.",
                            "El cambio supera el límite. Verifique el producto y tamaño del paquete."))
    elif not valid_price(current):
        points += weights["current"] * cfg.REVIEW_MISSING_CURRENT_CREDIT
        notes.append(tr("No positive current price is available; percentage change cannot be checked.",
                        "No hay precio actual positivo; no se puede verificar el cambio porcentual."))
    cv = number(row.get("Price CV"))
    if cv is not None:
        points += weights["spread"] * (1 if cv <= cv_limit else max(0, cv_limit / cv))
        notes.append(tr(f"Comparable-price variation: {cv:.0%} (limit {cv_limit:.0%}).",
                        f"Variación de precios comparables: {cv:.0%} (límite {cv_limit:.0%})."))
    else:
        notes.append(tr("Price variation is unavailable.", "La variación de precios no está disponible."))
    explanations = {
        "NO MATCHES": ("No comparable products survived the matching checks.", "No quedaron productos comparables."),
        "VERY HIGH RISK": ("Very little matching evidence supports this price.", "Muy poca evidencia respalda este precio."),
        "HIGH RISK": ("There are fewer comparable products than the evidence target.", "Hay menos comparables de los necesarios."),
        "ONE STORE ONLY": ("Only one supermarket supports this price.", "Solo un supermercado respalda este precio."),
        "LOW SEARCH CONFIDENCE": ("The search terms are broad or ambiguous.", "La búsqueda es amplia o ambigua."),
        "TARGET SIZE WARNING": ("The product name and spreadsheet disagree on package size.", "El nombre y la hoja indican tamaños distintos."),
        "REQUEST ERROR": ("A store request failed; evidence may be incomplete.", "Falló una consulta; la evidencia puede estar incompleta."),
    }
    flags = {part.strip().split(":", 1)[0] for part in str(row.get("Quality Flag") or "").split("|") if part.strip()} - {"OK"}
    if "PARTIAL REQUEST ERROR" in flags:
        flags.remove("PARTIAL REQUEST ERROR")
        flags.add("REQUEST ERROR")
    if row.get("Request Error"):
        flags.add("REQUEST ERROR")
    for flag in sorted(flags):
        points -= cfg.REVIEW_FLAG_PENALTIES.get(flag, cfg.REVIEW_UNKNOWN_FLAG_PENALTY)
        cap = min(cap, cfg.REVIEW_FLAG_CAPS.get(flag, 5))
        notes.append(tr(*explanations[flag]) if flag in explanations else tr("Warning: ", "Advertencia: ") + flag)
    for field in ("Search Warning", "Quality Flag", "Notes", "Target Quantity Source", "Request Error"):
        if row.get(field):
            notes.append(f"{field}: {row[field]}")
    for field in ("Size-Distant Removed", "Duplicates Removed", "Price Outliers Removed"):
        if number(row.get(field)):
            notes.append(f"{field}: {row[field]}")
    points = max(0, min(100, points))
    if not has_estimate:
        return 1, 0, notes
    return min(cap, 1 + sum(points >= boundary for boundary in cfg.REVIEW_SCORE_BOUNDARIES)), round(points, 1), notes


def reviewer_notes(item, settings, current_row=None, lang="en"):
    """Short observations for a human decision; full diagnostics stay in evidence."""
    tr = lambda en, es: es if lang == "es" else en
    row, current = item["row"], item["current"]
    count = int(number(row.get("Accepted Count")) or 0)
    stores = int(number(row.get("Store Count")) or len(item["counts"]))
    notes = [tr(f"Based on {count} comparable listings from {stores} stores.",
                f"Basado en {count} productos comparables de {stores} tiendas.")]
    if item["counts"]:
        notes[0] += " " + ", ".join(f"{name}: {n}" for name, n in sorted(item["counts"].items())) + "."
    if not valid_price(row.get(PRICE)):
        notes.append(tr("No usable estimate was found. A price can be entered manually.", "No se encontró una estimación válida. Puede ingresar un precio manual."))
    if not valid_price(current):
        notes.append(tr("No positive current price is available for comparison.", "No hay precio actual positivo para comparar."))
    else:
        proposed = number(row.get(PRICE))
        limit = settings.get("PRICE_CHANGE_REVIEW_THRESHOLD", cfg.PRICE_CHANGE_REVIEW_THRESHOLD)
        if proposed is not None and abs(proposed - current) / current > limit + 1e-12:
            notes.append(tr(f"The estimate differs from the current price by more than {limit:.0%}. Check the product and package size.",
                            f"La estimación difiere del precio actual en más de {limit:.0%}. Revise el producto y tamaño."))
    flags = str(row.get("Quality Flag") or "")
    if "TARGET SIZE WARNING" in flags:
        notes.append(tr("The product name and spreadsheet list different package sizes.", "El nombre y la hoja indican tamaños de paquete diferentes."))
    if str(row.get("Search Confidence") or "").upper() == "LOW":
        notes.append(tr("The search terms are broad; check whether the store products are comparable.", "La búsqueda es amplia; revise si los productos de tienda son comparables."))
    if row.get("Request Error") or "REQUEST ERROR" in flags:
        notes.append(tr("Some store searches failed, so the evidence may be incomplete.", "Fallaron consultas a tiendas; la evidencia puede estar incompleta."))
    variation = number(row.get("Price CV"))
    if variation is not None and variation > settings.get("PRICE_CV_WARNING", PRICE_CV_WARNING):
        notes.append(tr("Comparable store prices vary considerably.", "Los precios comparables varían considerablemente."))
    if not (current_row or {}).get("Id"):
        notes.append(tr("Salesforce Id is missing. This item needs a current-price record before CSV export.", "Falta el Id de Salesforce. Se necesita el registro de precio actual para exportar este producto."))
    return notes


# Column the robot adds to the results with each product's Salesforce PricebookEntry Id.
RESULTS_ID = "Salesforce Id"

# Columns of the current-prices file kept in the saved review (Id + price for the export, the rest for filters).
TEMPLATE_KEEP = (CODE, "ProductCode", "Id", "UnitPrice", "Precio de lista", NAME, "Sub-familia de Productos",
                 "Familia de productos", "Linea de producto", "Categoria RepTrim", "Tipo GFN")


class ReviewSession:
    def __init__(self, results, template=None, settings=None, products=None, lang="en"):
        # template (the Salesforce current-prices file) is optional: results made by this version carry the
        # Salesforce Id and current price of every product, and a saved review keeps its own copy.
        self.results = Path(results).resolve()
        self.template = Path(template).resolve() if template else None
        self.settings = dict(settings or {})
        self.lang = lang
        self.path = self.results.with_suffix(self.results.suffix + ".review.json")
        from candidate_store import file_hash
        results_hash = file_hash(self.results)
        saved = None
        if self.path.exists():
            try:
                saved = json.loads(self.path.read_text(encoding="utf-8"))
            except (ValueError, OSError):
                saved = None
        summary_rows = table(self.results, "Summary")
        if self.template is not None and self.template.is_file():
            self.fingerprint = [results_hash, file_hash(self.template)]
            self.template_rows = read_current_rows(self.template)
        elif saved and (saved.get("fingerprint") or [None])[0] == results_hash and saved.get("template_rows"):
            # The current-prices file was moved or deleted: use the copy saved with this review.
            self.fingerprint = list(saved["fingerprint"])
            self.template_rows = saved["template_rows"]
        elif summary_rows and RESULTS_ID in summary_rows[0]:
            # Results from this version: the Ids and current prices are in the results file itself.
            self.fingerprint = [results_hash, "results"]
            self.template_rows = [{CODE: code(r.get(CODE)), NAME: r.get(NAME), "Id": str(r.get(RESULTS_ID) or "").strip(),
                                   "Precio de lista": number(r.get("Current BAP Price"))}
                                  for r in summary_rows if code(r.get(CODE))]
        else:
            raise ValueError("This results file doesn't include the Salesforce Ids (it was made by an older version).\n"
                             "Choose the current prices file downloaded from Salesforce.")
        self.base = {}
        for row in self.template_rows:
            key = code(row.get(CODE))
            if key in self.base and key:
                raise ValueError(f"Duplicate product code in current prices: {key}")
            if key:
                self.base[key] = row
        self.product_rows = {}
        if products and Path(products).is_file():
            from price_robot import choose_sheet, find_header_row
            from openpyxl import load_workbook
            wb = load_workbook(products, read_only=True, data_only=True)
            try:
                ws = choose_sheet(wb)
                rows = ws.iter_rows(min_row=find_header_row(ws), values_only=True)
                headers = next(rows)
                self.product_rows = {code(r.get(CODE)): r for r in (dict(zip(headers, v)) for v in rows)}
            finally:
                wb.close()
        self.candidates = defaultdict(list)
        if self.results.suffix.lower() != ".csv":
            from candidate_store import CandidateStore
            self.candidates = CandidateStore(self.results, self.fingerprint[0])
        self.items = {}
        for row in summary_rows:
            key = code(row.get(CODE))
            if not key or key in self.items:
                raise ValueError(f"Missing or duplicate product code in results: {key}")
            counts = Counter(getattr(self.candidates, "counts", {}).get(key, {}))
            if not counts:
                counts.update({s.strip(): int(n) for s, n in re.findall(r"([^|=]+)=.*?\(n=(\d+)\)", str(row.get("Store Averages / kg") or ""))})
            current = number(self.base.get(key, {}).get("Precio de lista"))
            self.items[key] = {"row": row, "current": current, "counts": dict(counts),
                               "decision": "pending", "price": None, "note": "", "history": []}
        self.undo_stack = []
        self.redo_stack = []
        self.original_rows = {key: deepcopy(item["row"]) for key, item in self.items.items()}
        for item in self.items.values():
            item["candidate_overrides"] = {}
        self.archived_review = None
        if saved is not None:
            data = saved
            # Decisions belong to the results file; a different current-prices file doesn't undo them.
            if (data.get("fingerprint") or [None])[0] != self.fingerprint[0]:
                self.archived_review = self.path.with_name(self.path.name + datetime.now(timezone.utc).strftime(".stale-%Y%m%d-%H%M%S-%f"))
                self.path.rename(self.archived_review)
            else:
                self.settings = data.get("settings", self.settings)
                self.product_rows.update(data.get("product_rows", {}))
                for key, decision in data.get("decisions", {}).items():
                    if key in self.items:
                        self.items[key].update({k: decision[k] for k in ("decision", "price", "note", "history")})
                        self.items[key]["candidate_overrides"] = decision.get("candidate_overrides", {})
                        if "candidate_review" in decision:
                            reviewed = decision["candidate_review"]
                            self.items[key].update(row=reviewed["row"], counts=reviewed["counts"])
                            self.items[key]["candidate_review"] = reviewed
        for key, item in self.items.items():
            metadata = dict(self.base.get(key, {}))
            metadata.update({k: v for k, v in self.product_rows.get(key, {}).items() if v not in (None, "")})
            metadata.update({k: v for k, v in item["row"].items() if v not in (None, "")})
            item["categories"] = {field: str(metadata.get(field) or "").strip() for field in
                                  ("Sub-familia de Productos", "Familia de productos", "Linea de producto", "Categoria RepTrim", "Tipo GFN")}
            item["categories"]["Match group"] = str(item["row"].get("Generated Queries") or "").split("|")[0].strip()
        self.rescore()

    def rescore(self):
        for item in self.items.values():
            item["score"], item["points"], item["notes"] = confidence(item["row"], item["current"], item["counts"], self.settings, self.lang)

    def save(self):
        payload = {"fingerprint": self.fingerprint, "settings": self.settings,
                   "template": str(self.template or ""), "results": str(self.results),
                   # a copy of the Salesforce Ids and current prices, so the review and its export
                   # still work if the current-prices file is moved or deleted
                   "template_rows": [{k: row.get(k) for k in TEMPLATE_KEEP if row.get(k) not in (None, "")}
                                     for row in self.template_rows if code(row.get(CODE)) in self.items],
                   "product_rows": {key: self.product_rows[key] for key in self.items if key in self.product_rows},
                   "policy_version": cfg.REVIEW_POLICY_VERSION,
                   "decisions": {key: {k: item[k] for k in ("decision", "price", "note", "history", "candidate_overrides", "candidate_review") if k in item}
                                 for key, item in self.items.items()}}
        temp = self.path.with_suffix(self.path.suffix + ".tmp")
        temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        temp.replace(self.path)

    def candidate_accepted(self, key, index, overrides=None):
        overrides = self.items[key]["candidate_overrides"] if overrides is None else overrides
        return overrides.get(str(index), self.candidates[key][index].get("Status") == "ACCEPTED")

    @staticmethod
    def candidate_unit_price(candidate):
        value = number(candidate.get("Price / kg"))
        if value is not None and value > 0:
            return value
        price, quantity = number(candidate.get("Chosen Price")), number(candidate.get("Normalized kg"))
        if price is not None and price > 0 and quantity is not None and quantity > 0:
            return price / quantity
        return None

    def preview_candidates(self, key, overrides):
        """Recalculate from the human's selections using the scraper's price formula.

        Do not run matching/outlier filters again: the human explicitly chose
        these listings. Original candidate statuses and reasons remain intact.
        """
        import matcher_core as mc
        from price_robot import aggregate_stores
        tr = lambda en, es: es if self.lang == "es" else en
        candidates = self.candidates[key]
        if not candidates:
            raise ValueError(tr("No candidate details were saved for this product.",
                                "No se guardaron candidatos para este producto."))
        if any(not str(i).isdigit() or int(i) >= len(candidates) or type(value) is not bool
               for i, value in overrides.items()):
            raise ValueError("Invalid candidate selection")
        row = deepcopy(self.original_rows[key])
        quantity = number(row.get("Target Quantity kg"))
        if quantity is None or quantity <= 0:
            raise ValueError(tr("The target package quantity is missing. Enter a price manually.",
                                "Falta la cantidad del paquete objetivo. Ingrese un precio manual."))
        settings = dict(self.settings)
        for field in ("BAP_MARKET_PRICE_PERCENT", "STORE_WEIGHT_CAP", "ECONOMY_PERCENTILE"):
            settings[field] = number(row.get(field)) if number(row.get(field)) is not None else settings.get(field, getattr(mc, field))
        settings["MARKET_PRICE_LEVEL"] = row.get("Market Price Level") or settings.get("MARKET_PRICE_LEVEL", mc.MARKET_PRICE_LEVEL)
        stores = defaultdict(lambda: {"cleaned": []})
        for index, candidate in enumerate(candidates):
            if not self.candidate_accepted(key, index, overrides):
                continue
            value = self.candidate_unit_price(candidate)
            if value is None:
                raise ValueError(tr("Cannot include a listing without a usable price and size: ",
                                    "No se puede incluir un producto sin precio y tamaño válidos: ") + str(candidate.get("Candidate Product")))
            store = str(candidate.get("Store") or candidate.get("Store Name") or "Unknown")
            stores[store]["cleaned"].append({"price_per_kg": value})
        aggregate = aggregate_stores(stores, row.get("Average Mode") or "store_balanced", settings)
        count = len(aggregate["cleaned"])
        row.update({PRICE: aggregate["market"] * quantity * settings["BAP_MARKET_PRICE_PERCENT"] if aggregate["market"] is not None else None,
                    "Accepted Count": count, "Rejected Count": len(candidates) - count,
                    "Store Count": aggregate["store_count"], "Stores Used": ", ".join(aggregate["store_avgs"]),
                    "Store Averages / kg": " | ".join(f"{s}=${v:.2f}/kg (n={aggregate['store_counts'][s]})" for s, v in aggregate["store_avgs"].items())})
        for field, stat in (("Average Price / kg", "avg"), ("Median Price / kg", "median"),
                            ("Economy Price / kg", "economy"), ("Min Price / kg", "min"),
                            ("Max Price / kg", "max"), ("Price Ratio", "ratio"), ("Price CV", "cv")):
            row[field] = aggregate[stat]
        # Rebuild evidence-dependent warnings; retain request and target warnings.
        replaced = {"NO MATCHES", "VERY HIGH RISK", "HIGH RISK", "ONE STORE ONLY", "LOW SEARCH CONFIDENCE"}
        flags = [part.strip() for part in str(row.get("Quality Flag") or "").split("|")
                 if part.strip() and part.strip() != "OK" and part.strip().split(":", 1)[0] not in replaced]
        if not count:
            flags.append("NO MATCHES")
        elif count <= 2:
            flags.append(f"VERY HIGH RISK: only {count} selected match(es)")
        elif count < settings.get("MIN_GOOD_MATCHES", mc.MIN_GOOD_MATCHES):
            flags.append(f"HIGH RISK: only {count} selected match(es)")
        if aggregate["store_count"] == 1:
            flags.append("ONE STORE ONLY")
        if row.get("Search Confidence") == "LOW" and (count < settings.get("MIN_GOOD_MATCHES", mc.MIN_GOOD_MATCHES) or aggregate["store_count"] < 2):
            flags.append("LOW SEARCH CONFIDENCE")
        row["Quality Flag"] = " | ".join(flags) or "OK"
        return {"row": row, "counts": aggregate["store_counts"]}

    def apply_candidates(self, key, overrides):
        overrides = {str(i): value for i, value in overrides.items()}
        reviewed = self.preview_candidates(key, overrides)
        before = {key: deepcopy(self.items[key])}
        item = self.items[key]
        item.update(row=reviewed["row"], counts=reviewed["counts"], candidate_review=reviewed,
                    candidate_overrides=dict(overrides), decision="pending", price=None)
        item["history"].append({"at": datetime.now(timezone.utc).isoformat(), "action": "candidates",
                                "candidate_overrides": dict(overrides), "proposal": reviewed["row"].get(PRICE),
                                "decision": "pending"})
        try:
            self.rescore()
            self.save()
        except Exception:
            self.items.update(before)
            raise
        self.undo_stack.append(before)
        self.redo_stack.clear()   # a new change replaces anything that could be redone

    def decide(self, keys, decision, price=None, note=""):
        keys = list(keys)
        if decision not in ("accepted", "rejected", "manual", "pending"):
            raise ValueError("Unknown decision")
        if decision == "manual" and (len(keys) != 1 or not valid_price(price)):
            raise ValueError("Enter a positive price for one selected item.")
        # Accepting a product with no proposed price keeps its current BAP price.
        if decision == "accepted" and any(not valid_price(self.accept_price(k)) for k in keys):
            raise ValueError("Selection contains items with no proposed or current price. Enter those prices manually first.")
        before = {key: deepcopy(self.items[key]) for key in keys}
        for key in keys:
            item = self.items[key]
            value = price if decision == "manual" else self.accept_price(key) if decision == "accepted" else None
            item.update(decision=decision, price=round(number(value), cfg.REVIEW_PRICE_DECIMALS) if value is not None else None, note=note)
            item["history"].append({"at": datetime.now(timezone.utc).isoformat(), "decision": decision,
                                    "price": item["price"], "note": note})
        try:
            self.save()
        except Exception:
            self.items.update(before)
            raise
        self.undo_stack.append(before)
        self.redo_stack.clear()   # a new change replaces anything that could be redone

    def accept_price(self, key):
        """The price "Accept" approves: the robot's proposal, or the current BAP price when there is none."""
        item = self.items[key]
        proposal = item["row"].get(PRICE)
        return proposal if valid_price(proposal) else item["current"]

    def undo(self):
        if self.undo_stack:
            before = self.undo_stack[-1]
            after = {k: deepcopy(self.items[k]) for k in before}
            self.items.update(before)
            try:
                self.save()
            except Exception:
                self.items.update(after)
                raise
            self.undo_stack.pop()
            self.redo_stack.append(after)

    def redo(self):
        if self.redo_stack:
            state = self.redo_stack[-1]
            before = {k: deepcopy(self.items[k]) for k in state}
            self.items.update(deepcopy(state))
            try:
                self.rescore()
                self.save()
            except Exception:
                self.items.update(before)
                raise
            self.redo_stack.pop()
            self.undo_stack.append(before)

    def export(self, destination, include_unchanged=False):
        destination = Path(destination).resolve()
        if destination in (self.results, self.template, self.path):
            raise ValueError("Choose a new export filename; input files cannot be overwritten.")
        if destination.suffix.lower() != ".csv":
            raise ValueError("Choose a .csv filename for the Salesforce export.")
        # Ids and current prices are already loaded, so only a changed results file matters here.
        from candidate_store import file_hash
        if self.results.is_file() and file_hash(self.results) != self.fingerprint[0]:
            raise ValueError("The results file changed during review. Reopen the review before exporting.")
        prices = {key: item["price"] for key, item in self.items.items()
                  if item["decision"] in ("accepted", "manual")}
        if any(not valid_price(price) for price in prices.values()):
            raise ValueError("An approved price is invalid.")
        if include_unchanged:
            for key, row in self.base.items():
                if key not in prices and valid_price(row.get("Precio de lista")):
                    prices[key] = row["Precio de lista"]
        missing = [key for key in prices if not self.base.get(key, {}).get("Id")]
        if missing:
            raise ValueError("Cannot export: missing Salesforce PricebookEntry Id for " + ", ".join(sorted(missing)) +
                             ". Supply a current-prices file with ProductCode, Id and UnitPrice; product metadata cannot replace Id.")
        if not prices:
            raise ValueError("No approved prices to export.")
        self.save()
        temp = destination.with_name(destination.name + ".tmp")
        try:
            with temp.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["Id", "UnitPrice"])
                for key, price in prices.items():
                    writer.writerow([self.base[key]["Id"], f"{price:.{cfg.REVIEW_PRICE_DECIMALS}f}"])
            temp.replace(destination)
        finally:
            if temp.exists():
                temp.unlink()
        return len(prices)
