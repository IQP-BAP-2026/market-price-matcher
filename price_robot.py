from __future__ import annotations

import argparse
import csv
import math
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from statistics import mean, median, pstdev

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import matcher_core as mc
from store_parsers import PARSER_REGISTRY, available_stores, canonical_store_id


from app_runtime import DATA_DIR, OUTPUT_DIR, cache_file

BASE_DIR = DATA_DIR
PRODUCTS_FILE = BASE_DIR / "products.xlsx"
CACHE_FILE = cache_file("store_search_cache.sqlite3")
OUTPUT_BASENAME = "coincidencias_de_precios"
PERFORMANCE_BASENAME = str(OUTPUT_DIR / "analisis_de_desempeno.xlsx")
DEFAULT_WORKERS = 100
CACHE_TTL_HOURS = 12

# Optional callables, so the app can add to the output without re-opening large files afterwards:
#   SUMMARY_HOOKS:              hook(summary_headers, summary_rows) — add columns before anything is written
#   RESULTS_WORKBOOK_HOOKS:     hook(workbook, summary_headers, summary_rows) — just before the results are saved
#   PERFORMANCE_WORKBOOK_HOOKS: hook(workbook) — just before the performance analysis is saved
SUMMARY_HOOKS: list = []
RESULTS_WORKBOOK_HOOKS: list = []
PERFORMANCE_WORKBOOK_HOOKS: list = []


_ILLEGAL_XLSX_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def clean_cell(value):
    """Remove control characters that Excel files cannot store.

    Store product names occasionally contain invisible control characters
    (e.g. "SWEET GINGER CHI\x03 SAUCE"), which make openpyxl refuse to save.
    """
    if isinstance(value, str):
        return _ILLEGAL_XLSX_CHARS.sub("", value)
    return value


def normalize_product_code(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_money(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    try:
        return float(str(value).strip().replace("$", "").replace(",", ""))
    except (TypeError, ValueError):
        return None


def find_header_row(ws):
    wanted = {"nombre del producto", "codigo de producto"}
    for row in range(1, min(ws.max_row, 30) + 1):
        vals = {mc.normalize(ws.cell(row, c).value) for c in range(1, ws.max_column + 1)}
        if wanted.issubset(vals):
            return row
    raise ValueError("Could not find spreadsheet header row.")


def choose_sheet(wb):
    import re
    master = [s for s in wb.sheetnames if re.fullmatch(r"MASTER-\d{4}", s, re.I)]
    for name in master + wb.sheetnames:
        ws = wb[name]
        try:
            find_header_row(ws)
            return ws
        except ValueError:
            pass
    raise ValueError("No worksheet containing 'Nombre del producto' was found.")


def header_map(ws, header_row):
    return {
        mc.normalize(ws.cell(header_row, c).value): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(header_row, c).value is not None
    }


def style_sheet(ws):
    max_row, max_column = ws.max_row, ws.max_column
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, max_column + 1):
        width = min(48, max(12, max(
            len(str(ws.cell(r, col).value or "")) for r in range(1, min(max_row, 150) + 1)
        ) + 2))
        ws.column_dimensions[get_column_letter(col)].width = width


def candidate_from_store(raw: dict, target_name: str, spreadsheet_weight_kg, target_context: dict, query: str) -> dict:
    name = clean_cell(mc.repair_text(raw.get("name", ""))).strip()
    attrs = raw.get("attrs") or {}
    basis = mc.quantity_basis(target_name, target_context)
    ok, reason = mc.candidate_matches(target_name, name, attrs, target_context)
    qty = mc.parse_candidate_quantity(name, attrs, basis=basis)
    if ok and (qty.kg is None or qty.kg <= 0):
        ok = False
        reason = (
            "cannot count units in candidate title (target is priced per unit)"
            if basis == "count"
            else "cannot normalize candidate to kg from title or safe attributes"
        )

    size_ratio = target_unit_kg = candidate_unit_kg = None
    if ok:
        size_ratio, target_unit_kg, candidate_unit_kg = mc.package_size_ratio(
            target_name, name, spreadsheet_weight_kg, attrs, target_context
        )
        if size_ratio is None:
            candidate_unit_kg = qty.kg

    # Choose regular/final price with the CURRENT setting (cached search results
    # keep both prices, so a changed setting applies without refetching).
    regular_price = parse_money(raw.get("regular_price"))
    final_price = parse_money(raw.get("final_price"))
    if regular_price is not None or final_price is not None:
        price, price_basis = mc.choose_price(regular_price, final_price)
    else:
        price, price_basis = parse_money(raw.get("price")), raw.get("price_basis", mc.PRICE_BASIS)
    price_per_kg = (
        price / qty.kg
        if ok and price is not None and qty.kg is not None and qty.kg > 0
        else None
    )
    accepted = ok and price is not None and price_per_kg is not None
    if ok and price is None:
        reason = "missing usable retailer price"

    return {
        "store": raw.get("store", ""),
        "store_name": raw.get("store_name", raw.get("store", "")),
        "sku": str(raw.get("sku", "")).strip(),
        "name": name,
        "url": raw.get("url", ""),
        "query": query,
        "regular_price": regular_price,
        "final_price": final_price,
        "price": price,
        "price_basis": price_basis,
        "discount_percent": parse_money(raw.get("discount_percent")),
        "currency": raw.get("currency", "USD") or "USD",
        "kg": qty.kg,
        "target_unit_kg": target_unit_kg,
        "candidate_unit_kg": candidate_unit_kg,
        "size_ratio": size_ratio,
        "quantity_source": qty.source,
        # Meat sold by weight (deli counter) can be bought in any quantity: for
        # bulk meat boxes it joins the package-size group instead of choosing it.
        "size_neutral": qty.source.startswith("sold by weight")
                        and mc.normalize((target_context or {}).get("family")) in mc.SIZE_NEUTRAL_WEIGHED_FAMILIES,
        "accepted": accepted,
        "reason": reason,
        "price_per_kg": price_per_kg,
    }


def fetch_store_candidates(parser, product_name, spreadsheet_weight_kg, target_context, queries):
    seen = set()
    candidates = []
    error = ""

    def enough_raw():
        good = [c for c in candidates if c["accepted"]]
        close = [
            c for c in good
            if c.get("size_ratio") is not None and mc.MIN_SIZE_RATIO <= c["size_ratio"] <= mc.MAX_SIZE_RATIO
        ]
        return len(close) >= mc.MIN_GOOD_MATCHES or len(good) >= mc.RAW_MATCH_STOP_COUNT

    def take(page, query):
        for raw in page.get("items") or []:
            identity = (parser.store_id, str(raw.get("sku", "")).strip())
            if not identity[1] or identity in seen:
                continue
            seen.add(identity)
            candidates.append(candidate_from_store(
                raw, product_name, spreadsheet_weight_kg, target_context, query
            ))

    try:
        for query in queries:
            for page in parser.search_pages(query):
                take(page, query)
                if enough_raw():
                    break
            if enough_raw():
                break
        # Second pass for thin evidence: stores that support it are searched
        # a few pages deeper before settling for fewer than MIN_GOOD_MATCHES.
        deep = parser.deep_max_pages
        if deep > parser.max_pages and sum(1 for c in candidates if c["accepted"]) < mc.MIN_GOOD_MATCHES:
            for query in queries:
                for page in parser.search_pages(query, first_page=parser.max_pages + 1, last_page=deep):
                    take(page, query)
                    if enough_raw():
                        break
                if enough_raw():
                    break
    except Exception as exc:
        error = str(exc)

    return candidates, error


def clean_one_store(candidates: list[dict]):
    raw_accepted = [c for c in candidates if c["accepted"]]
    rejected = [c for c in candidates if not c["accepted"]]
    size_filtered, size_removed, size_strategy = mc.filter_preferred_size_matches(raw_accepted)
    deduped, duplicate_rows = mc.remove_near_duplicates(size_filtered)
    cleaned, outlier_rows = mc.remove_price_outliers(deduped)
    rejected.extend(size_removed)
    rejected.extend(duplicate_rows)
    rejected.extend(outlier_rows)
    return {
        "cleaned": cleaned,
        "rejected": rejected,
        "raw_accepted_count": len(raw_accepted),
        "size_removed_count": len(size_removed),
        "duplicate_count": len(duplicate_rows),
        "outlier_count": len(outlier_rows),
        "size_strategy": size_strategy,
    }


def clean_all_stores(store_candidates: dict[str, list[dict]]):
    """Clean the matches of all stores together.

    Package-size peers and price outliers are chosen across ALL stores, so a
    store that only sells small cans cannot pull the price per kg away from
    the package size the other stores agree on. Near-duplicates are removed
    within each store (the same product in two stores is two real prices).
    Returns the same per-store dictionaries as clean_one_store.
    """
    raw_accepted = [c for cands in store_candidates.values() for c in cands if c["accepted"]]
    size_kept, size_removed, size_strategy = mc.filter_preferred_size_matches(raw_accepted)

    deduped, duplicate_rows = [], []
    for store in store_candidates:
        kept, removed = mc.remove_near_duplicates([c for c in size_kept if c["store"] == store])
        deduped.extend(kept)
        duplicate_rows.extend(removed)
    cleaned, outlier_rows = mc.remove_price_outliers(deduped)

    out = {}
    for store, cands in store_candidates.items():
        mine = lambda rows: [c for c in rows if c["store"] == store]
        store_size = mine(size_removed)
        store_dups = mine(duplicate_rows)
        store_outliers = mine(outlier_rows)
        later = {id(c) for c in store_size + store_dups + store_outliers}
        out[store] = {
            "cleaned": mine(cleaned),
            "rejected": [c for c in cands if not c["accepted"] and id(c) not in later]
                        + store_size + store_dups + store_outliers,
            "raw_accepted_count": sum(1 for c in raw_accepted if c["store"] == store),
            "size_removed_count": len(store_size),
            "duplicate_count": len(store_dups),
            "outlier_count": len(store_outliers),
            "size_strategy": size_strategy,
        }
    return out


def aggregate_stores(store_results: dict[str, dict], average_mode: str, settings=None):
    """Combine the stores' clean matches into market prices per kg.

    Returns the average ("avg"), median, economy price and the market price
    ("market") for the MARKET_PRICE_LEVEL setting, which the estimate uses.
    """
    settings = settings or {}
    level = str(settings.get("MARKET_PRICE_LEVEL", mc.MARKET_PRICE_LEVEL) or "average").lower()
    economy_percentile = float(settings.get("ECONOMY_PERCENTILE", mc.ECONOMY_PERCENTILE))
    weight_cap = max(1, int(settings.get("STORE_WEIGHT_CAP", mc.STORE_WEIGHT_CAP)))
    cleaned_all = []
    store_avgs, store_counts, store_medians, store_economy = {}, {}, {}, {}
    for store, result in store_results.items():
        vals = [c["price_per_kg"] for c in result["cleaned"] if c.get("price_per_kg") is not None]
        if vals:
            store_avgs[store] = mean(vals)
            store_medians[store] = median(vals)
            store_economy[store] = mc.price_level_value(vals, "economy", economy_percentile)
            store_counts[store] = len(vals)
            cleaned_all.extend(result["cleaned"])

    if not cleaned_all:
        return {
            "avg": None, "median": None, "economy": None, "market": None, "level": level,
            "min": None, "max": None, "ratio": None, "cv": None, "store_avgs": store_avgs,
            "store_counts": store_counts, "store_medians": store_medians,
            "store_count": 0, "cleaned": [],
        }

    vals = [c["price_per_kg"] for c in cleaned_all]
    if average_mode == "store_balanced" and store_avgs:
        # Each store counts by its number of clean matches, capped, so no store
        # dominates and a single listing does not outweigh a store with several.
        weights = {s: min(store_counts[s], weight_cap) for s in store_avgs}

        def balanced(per_store: dict) -> float:
            return sum(per_store[s] * weights[s] for s in per_store) / sum(weights[s] for s in per_store)

        avg = balanced(store_avgs)
        economy = balanced(store_economy)
        market = {"average": avg, "economy": economy}.get(level, balanced(store_medians))
    else:
        avg = mean(vals)
        economy = mc.price_level_value(vals, "economy", economy_percentile)
        market = mc.price_level_value(vals, level, economy_percentile)
    med = median(vals)
    ratio = max(vals) / min(vals) if min(vals) > 0 else math.inf
    cv = pstdev(vals) / mean(vals) if len(vals) > 1 and mean(vals) else 0.0
    return {
        "avg": avg,
        "median": med,
        "economy": economy,
        "market": market,
        "level": level,
        "min": min(vals),
        "max": max(vals),
        "ratio": ratio,
        "cv": cv,
        "store_avgs": store_avgs,
        "store_counts": store_counts,
        "store_medians": store_medians,
        "store_count": len(store_avgs),
        "cleaned": cleaned_all,
    }


def process_product_multi(parsers, src, average_mode):
    queries, query_confidence, query_warning = mc.build_queries(src["name"], src.get("target_context"))
    per_store = {}
    request_errors = []

    store_candidates = {}
    for store_id, parser in parsers.items():
        candidates, error = fetch_store_candidates(
            parser, src["name"], src.get("weight_kg"), src.get("target_context") or {}, queries
        )
        for c in candidates:
            c["store"] = c.get("store") or store_id
        store_candidates[store_id] = candidates
        if error:
            request_errors.append(f"{store_id}: {error}")
    per_store = clean_all_stores(store_candidates)
    for store_id, candidates in store_candidates.items():
        per_store[store_id]["all_candidates"] = candidates

    agg = aggregate_stores(per_store, average_mode)
    target_qty = mc.target_quantity_kg(src["name"], src.get("weight_kg"), src.get("target_context"))
    estimated = (
        agg["market"] * target_qty.kg * mc.BAP_MARKET_PRICE_PERCENT
        if agg["market"] is not None and target_qty.kg is not None else None
    )

    total_clean = sum(len(x["cleaned"]) for x in per_store.values())
    total_rejected = sum(len(x["rejected"]) for x in per_store.values())
    flags = []
    if total_clean == 0:
        flags.append("NO MATCHES")
    elif total_clean <= 2:
        flags.append(f"VERY HIGH RISK: only {total_clean} clean match(es)")
    elif total_clean < mc.MIN_GOOD_MATCHES:
        flags.append(f"HIGH RISK: only {total_clean} clean match(es); need {mc.MIN_GOOD_MATCHES}")
    if agg["store_count"] == 1 and len(parsers) > 1:
        flags.append("ONE STORE ONLY")
    # A one-word search is only worth a flag when the evidence is also thin;
    # with several clean matches from 2+ stores the guardrails did their job.
    if query_confidence == "LOW" and (total_clean < mc.MIN_GOOD_MATCHES or agg["store_count"] < 2):
        flags.append("LOW SEARCH CONFIDENCE")
    ratio_limit, cv_limit = mc.PRICE_RATIO_WARNING, mc.PRICE_CV_WARNING
    if mc.is_variety_target(src["name"]):
        ratio_limit *= mc.VARIETY_VARIATION_FACTOR
        cv_limit *= mc.VARIETY_VARIATION_FACTOR
    # Informational notes: shown in the results but not a reason for review.
    # A wide price spread is normal (store brands vs premium) and the market
    # value level already accounts for it; in a 1,000-product test these
    # estimates matched current prices better than unflagged ones.
    notes = []
    if agg["ratio"] is not None and (mc.spread_ratio(
            [c["price_per_kg"] for c in agg["cleaned"]]) >= ratio_limit or (agg["cv"] or 0) >= cv_limit):
        notes.append(f"PRICE SPREAD: ratio={agg['ratio']:.2f}, CV={agg['cv']:.1%}")
    target_warning = mc.target_quantity_warning(src["name"], src.get("weight_kg"), src.get("target_context"))
    if target_warning:
        flags.append(target_warning)
    target_note = mc.target_quantity_note(src["name"], src.get("weight_kg"), src.get("target_context"))
    if target_note:
        notes.append(target_note)
    if request_errors:
        flags.append("PARTIAL REQUEST ERROR" if agg["store_count"] else "REQUEST ERROR")
    quality = " | ".join(flags) if flags else "OK"

    return {
        "queries": queries,
        "query_confidence": query_confidence,
        "query_warning": query_warning,
        "per_store": per_store,
        "aggregate": agg,
        "target_qty": target_qty,
        "estimated": estimated,
        "quality": quality,
        "notes": " | ".join(notes),
        "request_error": " | ".join(request_errors),
        "accepted_count": total_clean,
        "rejected_count": total_rejected,
    }


def compare_current_prices(summary_rows, summary_headers, current_prices_file="current_prices.xlsx"):
    # Intentionally local: do not replace this with a global path.
    current_path = Path(current_prices_file)
    if not current_path.exists():
        print(f"\n[PRICE COMPARISON SKIPPED] {current_path} not found.")
        return

    from current_prices import read_current_prices
    current_prices = read_current_prices(current_path)

    hi = {h: i for i, h in enumerate(summary_headers)}
    print("\n" + "=" * 105)
    print("CURRENT BAP PRICE vs MULTI-STORE ESTIMATE")
    print("=" * 105)
    compared = no_est = no_current = 0
    for row in summary_rows:
        code = normalize_product_code(row[hi["Código de producto"]])
        name = row[hi["Nombre del producto"]]
        current = current_prices.get(code)
        estimated = parse_money(row[hi["Estimated New Product Price"]])
        if not current:   # missing, or $0.00 in the current-prices file: nothing to compare
            no_current += 1
            continue
        if estimated is None:
            no_est += 1
            print(f"[NO ESTIMATE] {code} | {name}")
            continue
        min_ppkg = parse_money(row[hi["Min Price / kg"]])
        max_ppkg = parse_money(row[hi["Max Price / kg"]])
        target = parse_money(row[hi["Target Quantity kg"]])
        min_est = min_ppkg * target * mc.BAP_MARKET_PRICE_PERCENT if min_ppkg is not None and target else None
        max_est = max_ppkg * target * mc.BAP_MARKET_PRICE_PERCENT if max_ppkg is not None and target else None
        def pct(v):
            return "n/a" if v is None else f"{(v - current) / current * 100:+.1f}%"
        print(f"{code} | {name}")
        print(f"  Current ${current:.2f} | Estimate ${estimated:.2f} ({pct(estimated)})")
        if min_est is not None and max_est is not None:
            print(f"  Min ${min_est:.2f} ({pct(min_est)}) | Max ${max_est:.2f} ({pct(max_est)})")
        compared += 1
    print(f"Compared {compared}; no estimate {no_est}; no current price {no_current}")


def create_performance_analysis(source_rows, product_headers, summary_headers, summary_rows, detail_headers, detail_rows, current_prices_file="current_prices.xlsx"):
    # Intentionally local: do not replace this with a global path.
    current_path = Path(current_prices_file)
    current_prices = {}
    if current_path.exists():
        from current_prices import read_current_prices
        current_prices = read_current_prices(current_path)

    import results_format as rf
    out = BASE_DIR / PERFORMANCE_BASENAME
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = rf.DASHBOARD_SHEET
    analysis = wb.create_sheet(rf.ANALYSIS_SHEET)
    candidates = wb.create_sheet(rf.CANDIDATES_SHEET)

    # every product column, every robot column (translated), then the comparison and the review flag
    extra_headers = ([] if "Current BAP Price" in summary_headers else ["Current BAP Price", "% Difference vs Current"])
    extra_headers.append("Needs Manual Review")
    si = {h: i for i, h in enumerate(summary_headers)}
    summary_by_row = {r[si["Spreadsheet Row"]]: r for r in summary_rows}
    analysis_rows = []
    for src in source_rows:
        srow = summary_by_row.get(src["row"])
        if not srow:
            continue
        code = normalize_product_code(srow[si["Código de producto"]])
        current = current_prices.get(code)
        estimated = parse_money(srow[si["Estimated New Product Price"]])
        diff = None if current in (None, 0) or estimated is None else (estimated-current)/current
        quality = str(srow[si["Quality Flag"]] or "")
        needs_review = "NO" if quality == "OK" else "YES"
        original = list(src.get("original_values") or [])
        if len(original) < len(product_headers):
            original += [None] * (len(product_headers)-len(original))
        extras = ([] if "Current BAP Price" in summary_headers else [current, diff]) + [needs_review]
        analysis_rows.append(original[:len(product_headers)] + list(srow) + extras)
    # product columns keep their own (Spanish) names; the robot's columns are translated
    keys = [f"\x00product:{h}" for h in product_headers] + list(summary_headers) + extra_headers
    repeated = {"ERP Id (QBO)", "Código de producto", "Nombre del producto"} & set(product_headers)
    layout = [(k, True) for k in keys if k not in repeated] + [("=alerts", True)]   # no repeated product columns
    table = rf.Table(keys, analysis_rows, layout)
    table.headers = [str(k).split(":", 1)[1] if str(k).startswith("\x00product:") else h
                     for (k, _), h in zip(table.columns, table.headers)]
    rf.write_sheet(analysis, table, style_sheet)

    from excel_export import prepare_candidates, save_with_candidates
    detail_table = rf.full_table(detail_headers, detail_rows, kind="detail", extra_display=("=reason",))
    prepare_candidates(candidates, detail_table, style_sheet)

    total = len(summary_rows)
    est_idx = si["Estimated New Product Price"]
    store_idx = si["Store Count"]
    qual_idx = si["Quality Flag"]
    with_est = sum(1 for r in summary_rows if r[est_idx] is not None)
    multi_store = sum(1 for r in summary_rows if (r[store_idx] or 0) >= 2)
    ok = sum(1 for r in summary_rows if r[qual_idx] == "OK")
    rows = [
        ["Indicador", "Cantidad", "Porcentaje"],
        ["Productos procesados", total, 1 if total else 0],
        ["Productos con precio propuesto", with_est, with_est/total if total else 0],
        ["Productos con 2 o más supermercados", multi_store, multi_store/total if total else 0],
        ["Productos sin alertas", ok, ok/total if total else 0],
        ["Productos de supermercado revisados", len(detail_rows), None],
    ]
    for row in rows:
        dashboard.append(row)
    style_sheet(dashboard)
    for cell in dashboard[1]:
        cell.font = Font(color="FFFFFF", bold=True)
    for r in range(2, dashboard.max_row+1):
        dashboard.cell(r,3).number_format = '0.0%'
    for hook in PERFORMANCE_WORKBOOK_HOOKS:
        hook(wb)

    print(f"[EXPORT] Writing performance analysis: {len(detail_rows):,} candidates", flush=True)
    save_with_candidates(wb, out, detail_table,
                         lambda done, total: print(f"[EXPORT] Performance candidates: {done:,}/{total:,}", flush=True),
                         sheet=rf.CANDIDATES_SHEET)
    print(f"Performance analysis saved: {out}")


def write_results(output_path, summary_headers, summary_rows, detail_headers, detail_rows):
    """Save the results file. What BAP sees: the summary and the store listings, in Spanish. Columns that
    only the review reads are kept at the right, hidden; columns nobody needs are left out."""
    import results_format as rf
    output_path = Path(output_path)
    summary_table = rf.Table(summary_headers, summary_rows, rf.SUMMARY_LAYOUT)
    if output_path.suffix.lower() == ".csv":
        with output_path.open("w",newline="",encoding="utf-8-sig") as f:
            w=csv.writer(f); w.writerow(summary_table.headers); w.writerows(summary_table)
        return
    wb = Workbook(); sh=wb.active; sh.title=rf.SUMMARY_SHEET; dh=wb.create_sheet(rf.CANDIDATES_SHEET)
    rf.write_sheet(sh, summary_table, style_sheet)
    from excel_export import prepare_candidates, save_with_candidates
    candidates_table = rf.Table(detail_headers, detail_rows, rf.CANDIDATES_LAYOUT, kind="detail")
    prepare_candidates(dh, candidates_table, style_sheet)
    for hook in RESULTS_WORKBOOK_HOOKS:
        hook(wb, summary_headers, summary_rows)
    print(f"[EXPORT] Writing results: {len(detail_rows):,} candidates", flush=True)
    save_with_candidates(wb, output_path, candidates_table,
                         lambda done, total: print(f"[EXPORT] Results candidates: {done:,}/{total:,}", flush=True),
                         sheet=rf.CANDIDATES_SHEET)
    from candidate_store import build_candidate_index
    print("[EXPORT] Preparing fast review index", flush=True)
    try:
        build_candidate_index(output_path, detail_headers, detail_rows)
    except Exception as exc:
        print(f"[EXPORT] Review index will be rebuilt when opened: {exc}", flush=True)


def main():
    ap = argparse.ArgumentParser(description="BAP multi-store deterministic price robot")
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--contains", default=None)
    ap.add_argument("--type", dest="product_type", default=None)
    ap.add_argument("--list-types", action="store_true")
    ap.add_argument("--products", type=Path, default=PRODUCTS_FILE)
    ap.add_argument("--workers", type=int, default=DEFAULT_WORKERS)
    ap.add_argument("--stores", default="all", help="Comma-separated stores, e.g. super99,superxtra,rey,ribasmith (default: all)")
    ap.add_argument("--store", action="append", dest="store_repeat", help="Repeatable alias, e.g. --store super99 --store superxtra")
    ap.add_argument("--list-stores", action="store_true")
    ap.add_argument("--clear-cache", action="store_true")
    ap.add_argument("--no-cache", action="store_true")
    ap.add_argument("--average-mode", choices=["store_balanced", "all_candidates"], default="store_balanced")
    ap.add_argument("--output", default=str(OUTPUT_DIR / f"{OUTPUT_BASENAME}.xlsx"))
    args = ap.parse_args()

    if args.list_stores:
        print("Available stores:")
        for s in available_stores():
            print(f"  {s}")
        return

    if args.store_repeat:
        selected = [canonical_store_id(s) for s in args.store_repeat if s.strip()]
    elif args.stores.strip().lower() == "all":
        selected = available_stores()
    else:
        selected = [canonical_store_id(s) for s in args.stores.split(",") if s.strip()]
    bad = [s for s in selected if s not in PARSER_REGISTRY]
    if bad:
        raise ValueError(f"Unknown store(s): {', '.join(bad)}. Use --list-stores.")
    selected = list(dict.fromkeys(selected))

    input_path = args.products
    if not input_path.is_absolute():
        input_path = BASE_DIR / input_path
    if not input_path.exists():
        raise FileNotFoundError(f"Could not find product workbook: {input_path}")

    from table_files import load_book, to_number
    wb_in = load_book(input_path)   # .xlsx, .xls, .csv (any encoding) or an HTML "Excel" export
    ws = choose_sheet(wb_in)
    hr = find_header_row(ws)
    hm = header_map(ws, hr)
    from current_prices import PRODUCT_REQUIRED_COLUMNS
    missing = [c for c in PRODUCT_REQUIRED_COLUMNS if mc.normalize(c) not in hm]
    if missing:
        raise ValueError(f"The products sheet is missing required column(s): {', '.join(missing)}")
    name_col = hm["nombre del producto"]
    code_col = hm.get("codigo de producto")
    erp_col = hm.get("erp id qbo")
    weight_col = hm.get("unidad de peso en kg")
    type_col = hm.get("sub familia de productos")
    line_col = hm.get("linea de producto")
    family_col = hm.get("familia de productos")
    description_col = hm.get("descripcion del producto")
    group_col = hm.get("grupo de productos")
    category_col = hm.get("categoria reptrim")

    if args.list_types:
        vals = set()
        for row in ws.iter_rows(min_row=hr+1, values_only=True):
            v = row[type_col-1] if type_col else None
            if v not in (None, ""):
                vals.add(str(v).strip())
        for v in sorted(vals, key=mc.normalize):
            print(v)
        wb_in.close()
        return

    product_headers = [str(v).strip() if v is not None else f"Original Column {i}" for i,v in enumerate(next(ws.iter_rows(min_row=hr,max_row=hr,values_only=True)),1)]
    source_rows = []
    for r, values in enumerate(ws.iter_rows(min_row=hr+1, values_only=True), start=hr+1):
        namev = values[name_col-1]
        if not namev:
            continue
        name = str(namev).strip()
        if args.contains and mc.normalize(args.contains) not in mc.normalize(name):
            continue
        product_type = str(values[type_col-1]).strip() if type_col and values[type_col-1] is not None else ""
        if args.product_type and mc.normalize(args.product_type) != mc.normalize(product_type):
            continue
        source_rows.append({
            "row": r,
            "name": name,
            "code": values[code_col-1] if code_col else "",
            "erp": values[erp_col-1] if erp_col else "",
            # a text file gives "17" or "17,5": the matcher needs a number
            "weight_kg": to_number(values[weight_col-1], getattr(ws, "decimal_comma", False)) if weight_col else None,
            "product_type": product_type,
            "target_context": {
                "line": values[line_col-1] if line_col else "",
                "family": values[family_col-1] if family_col else "",
                "description": values[description_col-1] if description_col else "",
                "group": values[group_col-1] if group_col else "",
                "category": values[category_col-1] if category_col else "",
            },
            "original_values": list(values),
        })
        if args.limit and len(source_rows) >= args.limit:
            break
    wb_in.close()

    ttl = 0 if args.no_cache else CACHE_TTL_HOURS
    parsers = {
        s: PARSER_REGISTRY[s](cache_file=CACHE_FILE, cache_ttl_hours=ttl, pool_size=max(16,args.workers*4))
        for s in selected
    }
    if args.clear_cache:
        for parser in parsers.values():
            parser.clear_cache()
        print(f"Cleared cache for: {', '.join(selected)}")

    print(f"Stores: {', '.join(selected)}")
    print(f"Average mode: {args.average_mode}")
    print(f"Products: {len(source_rows)} | Workers: {max(1,args.workers)}")
    started = time.perf_counter()
    results = {}

    def run_one(src):
        return src["row"], process_product_multi(parsers, src, args.average_mode)

    try:
        with ThreadPoolExecutor(max_workers=max(1,args.workers)) as executor:
            futures = {executor.submit(run_one, src): src for src in source_rows}
            for i, future in enumerate(as_completed(futures),1):
                src = futures[future]
                try:
                    rowno, result = future.result()
                except Exception as exc:
                    rowno = src["row"]
                    result = {"error": str(exc)}
                results[rowno] = result
                print(f"[{i}/{len(source_rows)}] {src['name']}")
    finally:
        for parser in parsers.values():
            parser.close()

    summary_headers = [
        "Spreadsheet Row", "ERP Id (QBO)", "Código de producto", "Nombre del producto",
        "Generated Queries", "Search Confidence", "Search Warning", "Context Modifiers",
        "Selected Stores", "Stores Used", "Store Count", "Average Mode", "Store Averages / kg",
        "Raw Identity Matches", "Size-Distant Removed", "Duplicates Removed", "Price Outliers Removed",
        "Accepted Count", "Rejected Count", "Target Quantity kg", "Target Quantity Source",
        "Average Price / kg", "Median Price / kg", "Economy Price / kg", "Market Price Level",
        "Estimated New Product Price",
        "Min Price / kg", "Max Price / kg", "Price Ratio", "Price CV", "Price Basis",
        "Quality Flag", "Request Error", "Notes",
    ]
    detail_headers = [
        "Spreadsheet Row", "Código de producto", "Target Product", "Store", "Store Name", "Search Query",
        "Status", "SKU", "Candidate Product", "Product URL", "Regular Price", "Final Price", "Chosen Price",
        "Price Basis", "Discount %", "Currency", "Normalized kg", "Target Unit kg", "Candidate Unit kg",
        "Size Ratio", "Price / kg", "Quantity Parsed From", "Reason",
    ]
    summary_rows, detail_rows = [], []

    for src in source_rows:
        result = results.get(src["row"], {})
        if result.get("error"):
            queries, conf, warn = mc.build_queries(src["name"], src.get("target_context"))
            summary_rows.append([
                src["row"],src["erp"],src["code"],src["name"]," | ".join(queries),conf,warn,
                ", ".join(sorted(mc.context_modifier_tokens(src["name"],src.get("target_context")))),
                ", ".join(selected),"",0,args.average_mode,"",0,0,0,0,0,0,None,"",None,None,None,
                mc.MARKET_PRICE_LEVEL,None,None,None,None,None,mc.PRICE_BASIS,"REQUEST ERROR",result["error"],""
            ])
            continue

        agg = result["aggregate"]
        per_store = result["per_store"]
        stores_used = [s for s in selected if s in agg["store_avgs"]]
        raw_matches = sum(x["raw_accepted_count"] for x in per_store.values())
        size_removed = sum(x["size_removed_count"] for x in per_store.values())
        duplicates = sum(x["duplicate_count"] for x in per_store.values())
        outliers = sum(x["outlier_count"] for x in per_store.values())
        store_avgs_text = " | ".join(f"{s}=${agg['store_avgs'][s]:.2f}/kg (n={agg['store_counts'][s]})" for s in stores_used)
        summary_rows.append([
            src["row"],src["erp"],src["code"],src["name"]," | ".join(result["queries"]),
            result["query_confidence"],result["query_warning"],
            ", ".join(sorted(mc.context_modifier_tokens(src["name"],src.get("target_context")))),
            ", ".join(selected),", ".join(stores_used),agg["store_count"],args.average_mode,store_avgs_text,
            raw_matches,size_removed,duplicates,outliers,result["accepted_count"],result["rejected_count"],
            result["target_qty"].kg,result["target_qty"].source,
            agg["avg"],agg["median"],agg["economy"],agg["level"],
            result["estimated"],agg["min"],agg["max"],agg["ratio"],agg["cv"],mc.PRICE_BASIS,
            result["quality"],result["request_error"],result.get("notes", ""),
        ])
        for store_id in selected:
            store_result = per_store.get(store_id, {})
            cleaned_ids = {id(c) for c in store_result.get("cleaned", [])}
            for c in store_result.get("cleaned", []) + store_result.get("rejected", []):
                detail_rows.append([
                    src["row"],src["code"],src["name"],c.get("store"),c.get("store_name"),c.get("query"),
                    "ACCEPTED" if id(c) in cleaned_ids else "REJECTED",c.get("sku"),c.get("name"),c.get("url"),
                    c.get("regular_price"),c.get("final_price"),c.get("price"),c.get("price_basis"),c.get("discount_percent"),
                    c.get("currency"),c.get("kg"),c.get("target_unit_kg"),c.get("candidate_unit_kg"),c.get("size_ratio"),
                    c.get("price_per_kg"),c.get("quantity_source"),c.get("reason"),
                ])

    # Keep the calculation settings with the evidence for later human recalculation.
    calculation_fields = ("BAP_MARKET_PRICE_PERCENT", "STORE_WEIGHT_CAP", "ECONOMY_PERCENTILE")
    summary_headers.extend(calculation_fields)
    for row in summary_rows:
        row.extend(getattr(mc, field) for field in calculation_fields)

    for hook in SUMMARY_HOOKS:   # e.g. the current price and Salesforce Id of each product
        hook(summary_headers, summary_rows)

    # Last safety net before anything is written: no control characters in any cell.
    summary_rows = [[clean_cell(v) for v in row] for row in summary_rows]
    detail_rows = [[clean_cell(v) for v in row] for row in detail_rows]

    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = BASE_DIR / output_path
    write_results(output_path, summary_headers, summary_rows, detail_headers, detail_rows)

    elapsed=time.perf_counter()-started
    rate=len(source_rows)/elapsed*60 if elapsed and source_rows else 0
    print(f"\nSaved: {output_path}")
    print(f"Runtime: {elapsed/60:.1f} min ({rate:.1f} products/min)")
    compare_current_prices(summary_rows, summary_headers)
    create_performance_analysis(source_rows, product_headers, summary_headers, summary_rows, detail_headers, detail_rows)


if __name__ == "__main__":
    main()
